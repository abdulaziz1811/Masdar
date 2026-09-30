"""Tests for the Opendatasoft adapter against KAPSARC's recorded responses.

Recorded live on 2026-09-28. `catalog.json` is the real response to a search
for "electricity consumption"; the subscribers entry also carries the
fields, description, licence and themes returned by that dataset's own
metadata endpoint, since the catalogue call selected fewer columns.
`groupby_years.json` and `export.json` are verbatim; the export is the first
four records, so the fake reports a matching total_count of 4.
"""

import json
from datetime import date
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlsplit

import pytest
from openpyxl import load_workbook

from masdar.domain.models import Dimension, Verdict
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.adapters.opendatasoft import time_field
from masdar.sources.base import SourceError
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry, load_descriptors
from masdar.sources.transport import Response

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "kapsarc"
SUBSCRIBERS = "number-of-subscribers-by-branches-of-the-saudi-electricity-company-and-type-of-c"
TODAY = date(2026, 9, 28)


class KapsarcTransport:
    name = "recorded"

    def __init__(self, total_count=4):
        self.urls: list[str] = []
        self.total_count = total_count

    def get(self, url, source_id, params=None, headers=None):
        self.urls.append(url)
        path = urlsplit(url).path
        query = dict(parse_qsl(urlsplit(url).query))
        if path.endswith("/catalog/datasets"):
            # As the platform answers since September 2026: a language
            # variant named in the select is a malformed query.
            variants = [f.strip() for f in query.get("select", "").split(",")
                        if f.strip().endswith(("_ar", "_en"))]
            if variants:
                raise SourceError(source_id, f"HTTP 400: Unknown field: {variants[0]}")
            body = (FIXTURES / "catalog.json").read_bytes()
        elif path.endswith("/records") and "group_by" in query:
            body = (FIXTURES / "groupby_years.json").read_bytes()
        elif path.endswith("/records"):
            body = json.dumps({"total_count": self.total_count, "results": []}).encode()
        elif path.endswith("/exports/json"):
            body = (FIXTURES / "export.json").read_bytes()
        else:
            raise AssertionError(f"unexpected request: {url}")
        return Response(url, url, 200, body, "application/json", {})

    def describe(self):
        return self.name


def registry(transport):
    http = HttpClient(use_cache=False, retries=1, transport=transport)
    return Registry(tuple(d for d in load_descriptors() if d.id == "kapsarc"), http)


def agent(tmp_path, transport):
    reg = registry(transport)
    return Agent(registry=reg, http=reg._http,
                 config=AgentConfig(out_dir=tmp_path / "out", today=TODAY))


def subscribers(candidates):
    return next(c for c in candidates if c.dataset_id == SUBSCRIBERS)


class TestCatalogSearch:
    def test_searches_with_english_keywords(self):
        from masdar.nlu.parser import parse

        transport = KapsarcTransport()
        registry(transport).adapter("kapsarc").search(parse("استهلاك الكهرباء 2025"))
        where = dict(parse_qsl(urlsplit(transport.urls[0]).query))["where"]
        # Titles are English even in _ar fields; Arabic alone would miss them.
        assert '"Electricity"' in where and '"الطاقة الكهربائية"' in where
        assert " OR " in where

    def test_publisher_is_the_originating_body_credited_via_kapsarc(self):
        from masdar.nlu.parser import parse

        found = registry(KapsarcTransport()).adapter("kapsarc").search(parse("الكهرباء 2025"))
        candidate = subscribers(found)
        assert candidate.publisher_en.startswith("Saudi Central Bank (SAMA)")
        assert "KAPSARC" in candidate.publisher_en

    def test_last_update_is_the_publishers_own_date(self):
        from masdar.nlu.parser import parse

        found = registry(KapsarcTransport()).adapter("kapsarc").search(parse("الكهرباء 2025"))
        assert subscribers(found).last_updated == date(2026, 9, 15)

    def test_coverage_comes_from_the_data_not_the_description(self):
        from masdar.nlu.parser import parse

        found = registry(KapsarcTransport()).adapter("kapsarc").search(parse("الكهرباء 2025"))
        candidate = subscribers(found)
        # The description says "from 2005-2022"; the data runs to 2025.
        assert "2005-2022" in candidate.description
        assert candidate.claimed_coverage.years == frozenset(range(2005, 2026))
        assert candidate.claimed_coverage.origin.is_authoritative

    def test_coverage_groups_by_the_year_of_the_annotated_field(self):
        from masdar.nlu.parser import parse

        transport = KapsarcTransport()
        registry(transport).adapter("kapsarc").search(parse("الكهرباء 2025"))
        grouped = [unquote(u) for u in transport.urls if "group_by" in u]
        assert grouped and "group_by=year(year) as y" in grouped[0].replace("+", " ")

    def test_region_is_recognised_from_the_fields(self):
        from masdar.nlu.parser import parse

        found = registry(KapsarcTransport()).adapter("kapsarc").search(parse("الكهرباء 2025"))
        assert Dimension.REGION in subscribers(found).provided_dimensions


class TestTimeField:
    def test_annotation_beats_other_date_fields(self):
        fields = json.loads((FIXTURES / "catalog.json").read_text(encoding="utf-8"))
        subs = next(r for r in fields["results"] if r["dataset_id"] == SUBSCRIBERS)
        assert time_field(subs["fields"])["name"] == "year"

    def test_falls_back_to_a_field_named_year(self):
        assert time_field([{"name": "value", "type": "int"},
                           {"name": "year", "type": "int"}])["name"] == "year"

    def test_no_time_field_is_none(self):
        assert time_field([{"name": "value", "type": "int"}]) is None


class TestEndToEnd:
    def test_available_year_is_exported_with_its_real_update_date(self, tmp_path):
        answer = agent(tmp_path, KapsarcTransport()).answer(
            "ابي عدد مشتركي الكهرباء لسنه 2025 حسب المناطق"
        )
        assert answer.verdict is Verdict.AVAILABLE
        assert answer.primary.candidate.dataset_id == SUBSCRIBERS
        assert answer.primary.provenance.last_updated == date(2026, 9, 15)
        assert answer.primary.export_path
        text = "\n".join(
            str(c) for s in load_workbook(answer.primary.export_path).worksheets
            for row in s.iter_rows(values_only=True) for c in row if c is not None
        )
        assert "9307920" in text and "2026-09-15" in text
        assert "Saudi Central Bank (SAMA)" in text

    def test_unpublished_year_is_refused_with_the_latest(self, tmp_path):
        answer = agent(tmp_path, KapsarcTransport()).answer(
            "ابي عدد مشتركي الكهرباء لسنه 2026 حسب المناطق"
        )
        assert answer.verdict is Verdict.NOT_AVAILABLE
        assert [s.year for s in answer.suggestions] == [2025]

    def test_a_short_export_is_refused_not_delivered(self, tmp_path):
        # The platform says 735 records; the export returned 4.
        reg = registry(KapsarcTransport(total_count=735))
        url = reg.adapter("kapsarc")._url(f"/catalog/datasets/{SUBSCRIBERS}/exports/json")
        with pytest.raises(SourceError) as caught:
            reg.adapter("kapsarc").fetch(url)
        assert "735" in caught.value.reason


class TestNotAGastatSpec:
    def test_the_opendatasoft_spec_yields_no_cdata_datasets(self):
        from masdar.sources.openapi import datasets_from_spec

        spec = {"openapi": "3.0.3", "info": {"title": "Explore API"},
                "servers": [{"url": "https://datasource.kapsarc.org/api/explore/v2.1"}],
                "paths": {p: {"get": {"summary": "x"}} for p in (
                    "/catalog/datasets", "/catalog/exports/csv", "/catalog/facets",
                    "/catalog/datasets/{dataset_id}/records")}}
        # Read naively, these would be phantom datasets called "datasets",
        # "csv" and "facets".
        assert datasets_from_spec(spec, "https://api.stats.gov.sa") == []
