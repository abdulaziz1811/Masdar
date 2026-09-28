"""A source that cannot be reached is an outage, never evidence.

The catalogue of the GASTAT API is local (its OpenAPI specs), so its search
always "succeeds" -- even when the API itself is down. What follows must
still read as "could not reach", and must not cost the answer the sources
that did respond.
"""

from __future__ import annotations

from datetime import date
from urllib.parse import unquote

from masdar.domain.models import Verdict
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.base import SourceUnreachable
from masdar.sources.http import HttpClient
from masdar.sources.registry import DEMO_SOURCES_FILE, Registry, load_descriptors
from masdar.sources.transport import Transport
from tests.fakes import cdata_descriptor

QUESTION = "ابي احصاءات الطاقه الكهربائيه لسنه 2026 حسب المناطق"


class Down(Transport):
    name = "down"

    def __init__(self):
        self.datasets: list[str] = []
        self.calls = 0

    def get(self, url, source_id, params=None, headers=None):
        self.calls += 1
        # Count data downloads, not the year-only probes made while searching.
        decoded = unquote(url)
        if decoded.count("dimensions[]=") > 1:
            self.datasets.append(decoded.split("/v1/stats/")[-1].split("?")[0])
        raise SourceUnreachable(source_id, "انتهت مهلة الاتصال")


def agent_with_cdata_down(tmp_path):
    transport = Down()
    http = HttpClient(use_cache=False, retries=1, transport=transport)
    descriptors = (cdata_descriptor(), *load_descriptors(DEMO_SOURCES_FILE))
    agent = Agent(
        registry=Registry(descriptors, http),
        http=http,
        config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 9, 27)),
    )
    return agent, transport


def test_the_answer_rests_on_the_source_that_responded(tmp_path):
    agent, _ = agent_with_cdata_down(tmp_path)
    answer = agent.answer(QUESTION)
    assert answer.verdict is Verdict.NOT_AVAILABLE
    assert answer.primary.candidate.source_id != "gastat_cdata"
    assert [s.year for s in answer.suggestions] == [2024]


def test_the_outage_is_reported_as_one(tmp_path):
    agent, _ = agent_with_cdata_down(tmp_path)
    answer = agent.answer(QUESTION)
    assert "gastat_cdata" in {source for source, _ in answer.source_errors}
    cdata = [f for f in answer.findings if f.candidate.source_id == "gastat_cdata"]
    assert cdata and all(f.verdict is Verdict.SOURCE_UNREACHABLE for f in cdata)


def test_a_down_source_is_tried_once_per_question(tmp_path):
    # The host is asked once; after it fails to connect it is not asked again
    # for a short while (HttpClient's host cooldown), neither for the other
    # datasets' probes nor for their downloads.
    agent, transport = agent_with_cdata_down(tmp_path)
    agent.answer(QUESTION)
    assert transport.calls == 1


def test_without_the_cooldown_each_dataset_is_still_tried_at_most_once(tmp_path, monkeypatch):
    monkeypatch.setenv("MASDAR_HOST_COOLDOWN_SECONDS", "0")
    agent, transport = agent_with_cdata_down(tmp_path)
    agent.answer(QUESTION)
    assert len(set(transport.datasets)) == 1


def test_everything_down_says_so(tmp_path):
    transport = Down()
    http = HttpClient(use_cache=False, retries=1, transport=transport)
    agent = Agent(
        registry=Registry((cdata_descriptor(),), http),
        http=http,
        config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 9, 27)),
    )
    answer = agent.answer(QUESTION)
    assert answer.verdict is Verdict.SOURCE_UNREACHABLE
    assert "لا يعني أن البيانات غير موجودة" in answer.message_ar


def test_a_saved_copy_answers_when_the_source_is_down_and_says_so(tmp_path, monkeypatch):
    """During a presentation a ministry may be down; the last saved copy then
    answers, labelled with its date and the warning that newer data may exist."""
    import json
    from datetime import UTC, datetime, timedelta

    from openpyxl import load_workbook

    from tests.test_worldbank import Recorded, worldbank

    monkeypatch.setenv("MASDAR_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.delenv("MASDAR_STALE_MAX_DAYS", raising=False)
    question = "استهلاك الكهرباء للفرد 2023"

    def agent(transport):
        http = HttpClient(retries=1, transport=transport)
        return Agent(registry=Registry((worldbank(),), http), http=http,
                     config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 9, 28)))

    live = agent(Recorded())
    assert live.answer(question).verdict is Verdict.AVAILABLE
    for meta_path in (tmp_path / "cache").rglob("*.json"):  # two days pass
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        meta["retrieved_at"] = (datetime.now(UTC) - timedelta(days=2)).isoformat()
        meta_path.write_text(json.dumps(meta), encoding="utf-8")

    answer = agent(Down()).answer(question)
    assert answer.verdict is Verdict.AVAILABLE
    assert answer.primary.provenance.stale
    assert any("آخر نسخة محفوظة" in n and "أحدث" in n for n in answer.primary.notes)
    from masdar.export.excel import SOURCE_SHEET

    sheet = load_workbook(answer.primary.export_path)[SOURCE_SHEET]
    assert any("نسخة محفوظة" in str(c) for row in sheet.iter_rows(values_only=True) for c in row)
