"""Live indicators: HRSD's real-time labour figures and MEWA's counters.

Recorded through a Saudi exit on 2026-09-28. The point under test is what a
live figure may be said to be: a snapshot stamped with its retrieval date,
available for that year only, never passed off as an annual statistic.
"""

from __future__ import annotations

import hashlib
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from masdar.domain.models import Verdict
from masdar.nlu.parser import parse
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.base import SourceError
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry, load_descriptors
from masdar.sources.transport import Response, Transport

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "live"
ROUTES = {
    "node/5557175": "hrsd_labour.json",
    "getCamelsCount": "mewa_camels.json",
}


class Recorded(Transport):
    name = "recorded"

    def __init__(self, overrides=None):
        self.routes = {**ROUTES, **(overrides or {})}

    def get(self, url, source_id, params=None, headers=None):
        for fragment, name in self.routes.items():
            if fragment in url:
                body = (FIXTURES / name).read_bytes()
                return Response(url, url, 200, body, "application/json", {})
        raise AssertionError(f"unexpected request: {url}")


def registry(transport=None):
    http = HttpClient(use_cache=False, retries=1, transport=transport or Recorded())
    live = tuple(d for d in load_descriptors() if d.id in ("hrsd_live", "mewa_live"))
    assert len(live) == 2, "both live sources must be registered"
    return Registry(live, http)


def agent(tmp_path, transport=None):
    reg = registry(transport)
    return Agent(registry=reg, http=reg._http,
                 config=AgentConfig(out_dir=tmp_path / "out", today=date.today()))


class TestSearch:
    def test_a_workforce_question_finds_the_labour_indicators(self):
        found = registry().adapter("hrsd_live").search(parse("عدد العمالة السعوديين"))
        assert found[0].dataset_id == "labour-market-indicators"
        assert found[0].publisher_ar == "وزارة الموارد البشرية والتنمية الاجتماعية"

    def test_a_camel_question_finds_the_counter(self):
        found = registry().adapter("mewa_live").search(parse("كم عدد الإبل في المملكة"))
        assert [c.dataset_id for c in found][:1] == ["camels"]

    def test_unrelated_questions_find_nothing(self):
        assert registry().adapter("mewa_live").search(parse("الناتج المحلي 2022")) == []


class TestReshaping:
    def test_records_are_stamped_with_the_retrieval_date(self):
        adapter = registry().adapter("hrsd_live")
        url = adapter._indicators()[0]["url"]
        fetched = adapter.fetch(url)
        rows = json.loads(fetched.content)
        assert rows[0]["السنة"] == fetched.retrieved_at.year
        assert rows[0]["تاريخ الاستخراج"] == fetched.retrieved_at.date().isoformat()
        assert rows[0]["القيمة"] == 12618765

    def test_the_fingerprint_is_of_what_the_source_sent(self):
        adapter = registry().adapter("hrsd_live")
        fetched = adapter.fetch(adapter._indicators()[0]["url"])
        original = (FIXTURES / "hrsd_labour.json").read_bytes()
        assert fetched.sha256 == hashlib.sha256(original).hexdigest()

    def test_the_title_says_what_the_ministry_counts(self):
        # The ministry's own label is «عدد الإبل المرقمة في المملكة حالياً»:
        # tagged camels, which a reader must not take for every camel.
        (found,) = registry().adapter("mewa_live").search(parse("كم عدد الإبل في المملكة"))
        assert "المرقّمة" in found.title_ar
        assert "لا كل الإبل" in found.caveats[0]

    def test_a_counter_becomes_one_labelled_row(self):
        adapter = registry().adapter("mewa_live")
        camels = next(i for i in adapter._indicators() if i["id"] == "camels")
        rows = json.loads(adapter.fetch(camels["url"]).content)
        assert rows == [{
            "السنة": rows[0]["السنة"], "تاريخ الاستخراج": rows[0]["تاريخ الاستخراج"],
            "المؤشر": "عدد الإبل المرقّمة", "القيمة": 1924572,
        }]

    def test_an_empty_counter_is_an_error_not_a_zero(self):
        transport = Recorded({"getCamelsCount": "mewa_wells.json"})
        adapter = registry(transport).adapter("mewa_live")
        camels = next(i for i in adapter._indicators() if i["id"] == "camels")
        with pytest.raises(SourceError, match="لم يُرجع قيمة"):
            adapter.fetch(camels["url"])


class TestAnswers:
    def test_the_latest_figure_is_available_and_exported(self, tmp_path):
        answer = agent(tmp_path).answer("عدد الإبل")
        assert answer.verdict is Verdict.AVAILABLE
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        values = [c for row in sheet.iter_rows(values_only=True) for c in row]
        assert 1924572 in values

    def test_the_answer_says_it_is_a_snapshot(self, tmp_path):
        answer = agent(tmp_path).answer("عدد الإبل")
        assert any("بيانات لحظية" in note for note in answer.primary.notes)

    def test_an_earlier_year_is_not_claimed(self, tmp_path):
        answer = agent(tmp_path).answer("عدد الإبل 2022")
        assert answer.verdict is Verdict.NOT_AVAILABLE
        assert [s.year for s in answer.suggestions] == [date.today().year]
