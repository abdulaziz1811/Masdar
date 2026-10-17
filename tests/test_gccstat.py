"""GCC-Stat's SDMX service: the Kingdom's figures as the Gulf centre holds them.

The fixtures are the service's real answers for the Kingdom
(`.../data/GCCSTAT.MOBILE,<flow>,1.0/SAU`), recorded on 2026-10-04 and kept
byte for byte apart from whitespace. Two of them were asked from 2018 on
(`startPeriod`), which the service also answers.
"""

from __future__ import annotations

import json
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from masdar.chat.overview import overview
from masdar.domain.models import Verdict
from masdar.nlu.lexicon import load_lexicon
from masdar.nlu.parser import parse
from masdar.pipeline import resolve
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.adapters.gccstat import mark_repeats, sdmx_rows
from masdar.sources.base import SourceError
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry, load_descriptors
from masdar.sources.transport import Response, Transport

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "gccstat"


class Recorded(Transport):
    name = "recorded"

    def __init__(self):
        self.urls: list[str] = []

    def get(self, url, source_id, params=None, headers=None):
        self.urls.append(url)
        flow = url.split("GCCSTAT.MOBILE,")[1].split(",")[0]
        path = FIXTURES / f"{flow}.json"
        if not path.exists():
            raise AssertionError(f"unexpected request: {url}")
        return Response(url, url, 200, path.read_bytes(), "application/json", {})


def gccstat():
    return next(d for d in load_descriptors() if d.id == "gccstat")


def registry(transport=None):
    http = HttpClient(use_cache=False, retries=1, transport=transport or Recorded())
    return Registry((gccstat(),), http)


def recorded(flow: str) -> dict:
    return json.loads((FIXTURES / f"{flow}.json").read_text(encoding="utf-8"))


class TestConfiguration:
    def test_a_verified_regional_source_ranked_as_international(self):
        descriptor = gccstat()
        assert descriptor.api_verified and descriptor.enabled
        assert descriptor.international and descriptor.regional
        assert not descriptor.geo_restricted

    def test_every_indicator_names_its_unit(self):
        indicators = registry().adapter("gccstat")._indicators()
        assert len(indicators) >= 20
        for indicator in indicators:
            assert indicator["id"].startswith("MOBILE_DF_")
            assert indicator["title_ar"] and indicator["unit_ar"], indicator["id"]
            assert set(indicator["topics"]) <= set(gccstat().topics)

    def test_the_landing_page_counts_it_as_regional(self):
        reg = registry()
        listed = {s["id"]: s for s in overview(reg, load_lexicon())["sources"]}
        assert listed["gccstat"]["kind"] == "إقليمي"
        assert listed["gccstat"]["count"] == len(reg.adapter("gccstat")._indicators())


class TestSearch:
    @pytest.mark.parametrize("question, flow", [
        ("عدد السكان 2024", "MOBILE_DF_POP"),
        ("عدد سكان السعودية", "MOBILE_DF_POP"),
        ("معدل البطالة", "MOBILE_DF_SDG08_8_5_2"),
        ("الناتج المحلي الاجمالي 2023", "MOBILE_DF_NA_GDP_CUR"),
        ("استهلاك الكهرباء 2024", "MOBILE_DF_ENR_ELECT_CONS"),
        ("عدد السياح 2024", "MOBILE_DF_TUR_TOURISTS"),
        ("السياحة الوافدة", "MOBILE_DF_TUR_TOURISTS"),
        ("عدد الوفيات", "MOBILE_DF_DTH"),
        ("عدد المواليد 2022", "MOBILE_DF_BTH"),
        ("إنتاج التمور", "MOBILE_DF_AGR_Dates"),
        ("كمية المياه المحلاة", "MOBILE_DF_WTR_DESALINATED_WATER"),
        ("معدل التضخم 2025", "MOBILE_DF_CPI_INF_Y"),
    ])
    def test_the_indicator_asked_for_comes_first(self, question, flow):
        found = registry().adapter("gccstat").search(parse(question))
        assert found and found[0].dataset_id == flow

    @pytest.mark.parametrize("question", ["كم صيدلية في الرياض", "عدد الحجاج 2023",
                                          "أسعار الأسمنت"])
    def test_what_it_does_not_hold_finds_nothing(self, question):
        assert registry().adapter("gccstat").search(parse(question)) == []

    def test_a_candidate_names_the_centre_the_unit_and_the_flow(self):
        candidate = registry().adapter("gccstat").search(parse("الناتج المحلي الاجمالي"))[0]
        assert candidate.publisher_ar.startswith("المركز الإحصائي")
        assert candidate.caveats[0].startswith("مصدر إقليمي")
        assert "الوحدة: مليار دولار أمريكي." in candidate.caveats
        assert any("MOBILE_DF_NA_GDP_CUR" in c for c in candidate.caveats)
        assert candidate.landing_url.startswith("https://dp.marsa.gccstat.org")


class TestReadingSdmx:
    def test_one_row_per_year_with_the_unit(self):
        rows = sdmx_rows(recorded("MOBILE_DF_NA_GDP_CUR"), "الناتج", "مليار دولار أمريكي")
        assert [r["السنة"] for r in rows] == list(range(2010, 2026))
        assert rows[-1] == {"السنة": 2025, "المؤشر": "الناتج", "القيمة": 1273.6,
                            "الوحدة": "مليار دولار أمريكي", "ملاحظة": ""}

    def test_a_whole_number_stays_whole(self):
        rows = sdmx_rows(recorded("MOBILE_DF_POP"), "عدد السكان", "نسمة")
        assert rows[-1]["السنة"] == 2024 and rows[-1]["القيمة"] == 35300280
        assert isinstance(rows[-1]["القيمة"], int)

    def test_the_centres_arabic_footnote_is_kept_on_its_year(self):
        rows = {r["السنة"]: r for r in sdmx_rows(recorded("MOBILE_DF_POP"), "x", "")}
        assert rows[2023]["ملاحظة"] == "السعودية: بيان 2023م مرحل من عام 2022م"
        assert rows[2024]["ملاحظة"] == ""

    def test_a_note_about_other_member_states_is_left_out(self):
        # «بيانات الإمارات والبحرين (2021-2023) من منظمة العمل الدولية»
        rows = sdmx_rows(recorded("MOBILE_DF_LAB_EMPLOYEES"), "x", "")
        assert rows[-1]["السنة"] == 2025 and rows[-1]["القيمة"] == 19034128
        assert all(r["ملاحظة"] == "" for r in rows)

    def test_another_shape_is_an_error(self):
        with pytest.raises(SourceError, match="البنية"):
            sdmx_rows({"data": {}}, "x", "")


class TestRepeatedValues:
    def test_deaths_repeated_for_five_years_are_flagged(self):
        rows = sdmx_rows(recorded("MOBILE_DF_DTH"), "عدد الوفيات", "حالة وفاة")
        assert [r["القيمة"] for r in rows] == [80246] * 5
        note = mark_repeats(rows)
        assert "80,246" in note and "2018–2022" in note
        assert rows[0]["ملاحظة"] == ""
        assert all("مرحّلة" in r["ملاحظة"] and "2018" in r["ملاحظة"] for r in rows[1:])

    def test_births_repeated_after_a_change_are_flagged_from_the_change(self):
        rows = sdmx_rows(recorded("MOBILE_DF_BTH"), "عدد المواليد", "مولود")
        note = mark_repeats(rows)
        assert "2019–2022" in note
        by_year = {r["السنة"]: r["ملاحظة"] for r in rows}
        assert by_year[2018] == by_year[2019] == ""
        assert "2019" in by_year[2022]

    def test_two_equal_years_are_not_a_run(self):
        rows = sdmx_rows(recorded("MOBILE_DF_SDG08_8_5_2"),
                         "x", "")
        assert mark_repeats(rows) == ""

    def test_a_small_rate_repeating_is_not_judged(self):
        rows = [{"السنة": y, "القيمة": 2.5, "ملاحظة": ""} for y in (2021, 2022, 2023)]
        assert mark_repeats(rows) == ""


class TestAnswers:
    def agent(self, tmp_path, transport=None):
        reg = registry(transport)
        return Agent(registry=reg, http=reg._http,
                     config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 10, 4)))

    def test_a_published_year_is_available_with_its_unit(self, tmp_path):
        answer = self.agent(tmp_path).answer("عدد السكان 2024")
        assert answer.verdict is Verdict.AVAILABLE
        primary = answer.primary
        assert primary.candidate.dataset_id == "MOBILE_DF_POP"
        sheet = load_workbook(primary.export_path)["البيانات"]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        assert rows[0][:4] == ["السنة", "المؤشر", "القيمة", "الوحدة"]
        assert [2024, "عدد السكان", 35300280, "نسمة"] in [r[:4] for r in rows]

    def test_the_answer_says_it_is_regional(self, tmp_path):
        answer = self.agent(tmp_path).answer("الناتج المحلي الاجمالي 2025")
        assert answer.verdict is Verdict.AVAILABLE
        assert any(n.startswith("مصدر إقليمي") for n in answer.primary.notes)
        assert "مليار دولار أمريكي" in " ".join(answer.primary.notes)

    def test_a_carried_forward_value_is_said_to_be_one(self, tmp_path):
        answer = self.agent(tmp_path).answer("عدد الوفيات 2022")
        notes = " ".join(answer.primary.notes)
        assert "مرحّلة" in notes and "80,246" in notes

    def test_a_year_after_the_series_is_not_available(self, tmp_path):
        answer = self.agent(tmp_path).answer("عدد الوفيات 2024")
        assert answer.verdict is Verdict.NOT_AVAILABLE
        assert 2022 in [s.year for s in answer.suggestions]


class TestSaudiFirst:
    def test_a_regional_copy_scores_below_the_same_saudi_result(self):
        request = parse("عدد السكان")
        candidate = registry().adapter("gccstat").search(request)[0]
        saudi = replace(gccstat(), id="saudi", international=False, regional=False)
        local = resolve.score_candidate(request, replace(candidate, source_id="saudi"), saudi)
        regional = resolve.score_candidate(request, candidate, gccstat())
        assert regional.score < local.score
        assert any(r.startswith("مصدر إقليمي") for r in regional.match_reasons)

    def test_topics_ask_it_before_the_world_bank(self):
        lexicon = load_lexicon()
        for topic in lexicon.topics:
            preferred = list(topic.preferred_sources)
            if "gccstat" in preferred and "worldbank" in preferred:
                assert preferred.index("gccstat") < preferred.index("worldbank"), topic.id
