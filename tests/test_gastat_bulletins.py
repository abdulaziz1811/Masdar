"""GASTAT's bulletin workbooks, read as a statistician would read them.

`fixtures/gastat_site/workbooks/*.rebuilt.xlsx` are GASTAT's own workbooks
(labour market Q4 2025 and Q2 2026, register-based labour market Q4 2025,
CPI August 2026, GDP Q2 2026), read through Firecrawl on 2026-10-04 and
rebuilt sheet by sheet with their cell values verbatim; merged cells and
styling are not reproduced, and each sheet is cut after its first rows to
keep the files small. The search pages are those of
test_gastat_questions.py; the bulletin pages are reduced to the one
/documents/ link each carries, as recorded.
"""

from __future__ import annotations

import gzip
import json
from datetime import date
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlsplit

import pytest
from openpyxl import load_workbook

from masdar.domain.models import CoverageOrigin, Verdict
from masdar.export.tabular import YearAxis, read_xlsx
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry, load_descriptors
from masdar.sources.transport import Response

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "gastat_site"
BOOKS = FIXTURES / "workbooks"
SEARCHES = FIXTURES / "questions"
QUERIES = json.loads((SEARCHES / "queries.json").read_text(encoding="utf-8"))


def book(name: str) -> bytes:
    return (BOOKS / f"{name}.rebuilt.xlsx").read_bytes()


class TestTheTableAQuestionIsAnsweredBy:
    @pytest.mark.parametrize("workbook, question, sheet", [
        ("labour_q2_2026", "معدل البطالة 2025", "2-1"),
        ("labour_q2_2026", "معدل البطالة حسب المنطقة", "2-4"),
        ("labour_q2_2026", "معدل المشاركة في القوى العاملة", "5-1"),
        ("labour_q2_2026", "متوسط الأجر الشهري حسب النشاط الاقتصادي", "6-3"),
        # Nothing narrower asked: GASTAT's headline table, not the longest.
        ("labour_q2_2026", "أحدث بيانات سوق العمل", "1"),
        ("cpi_aug_2026", "معدل التضخم أغسطس 2026", "1.1"),
        ("cpi_aug_2026", "أحدث مؤشر لأسعار المستهلك", "1.1"),
        ("gdp_q2_2026", "الناتج المحلي الإجمالي 2026", "1.1"),
        ("gdp_q2_2026", "نمو الناتج المحلي الحقيقي", "2.3"),
    ])
    def test_the_right_sheet(self, workbook, question, sheet):
        assert read_xlsx(book(workbook), question).sheet_name == sheet

    @pytest.mark.parametrize("workbook", ["cpi_aug_2026", "gdp_q2_2026", "register_q4_2025"])
    def test_a_contents_or_introduction_sheet_is_never_the_data(self, workbook):
        # «المحتوى | Content», «القائمة الرئيسية», «المقدمة»: each was once
        # chosen over every table of its workbook.
        for question in ("", "الناتج المحلي الإجمالي", "عدد المشتغلين 2025", "أحدث البيانات"):
            table = read_xlsx(book(workbook), question)
            assert table.has_figures and not table.is_contents, (workbook, question)


class TestReadingGastatsTables:
    def table(self):
        return read_xlsx(book("labour_q2_2026"), "معدل البطالة 2025")

    def test_a_series_of_quarters_is_a_year_axis(self):
        # The source and method notes below the series kept it under the
        # share a period column needs: it was read as 2026 alone.
        table = self.table()
        assert table.year_axis() is YearAxis.ROWS
        assert table.observed_years() == frozenset(range(2021, 2027))

    def test_a_two_row_header_names_every_column(self):
        columns = self.table().columns
        assert columns[0] == "الربع / السنة"
        assert "السعوديين - الذكور" in columns and "غير السعوديين - الإجمالي" in columns

    def test_a_note_naming_the_year_is_not_a_row_of_it(self):
        rows = self.table().filter_years(frozenset({2025})).rows
        assert [r[0] for r in rows] == ["الربع الأول 2025", "الربع الثاني 2025",
                                        "الربع الثالث 2025", "الربع الرابع 2025", 2025]


# -- end to end ------------------------------------------------------------

BULLETINS = {
    "/ar/w/إحصاءات-سوق-العمل-للربع-الرابع-من-عام-2025م":
        "/documents/20117/2435273/Labor_Market_Statistics_Q4_2025_AR_%281%29.xlsx_fixed_13014526/"
        "880e1667-8d01-231d-0529-32a9a0fc9176?t=1774932582746",
    "/ar/w/الإحصاءات-السجليَّة-لسوق-العمل-للربع-الرابع-لعا-5":
        "/documents/20117/2435273/Register-based_Labour_Market_Statistics-_Q4_2025_AR.xlsx_fixed_"
        "13014495/d0ae4cc6-1a99-8665-21f3-c9aa4cf7dead?t=1774932482879",
    "/ar/w/labor-market-statistics-q2-of-2026":
        "/documents/20117/2435273/Labor+Market+Statistics+Q2+2026+AR+%282%29.xlsx/"
        "5c6df9fc-d491-598a-6fd3-acc335267279?t=1790766243679",
}
WORKBOOKS = {
    "Labor_Market_Statistics_Q4_2025": "labour_q4_2025",
    "Register-based_Labour_Market_Statistics-_Q4_2025": "register_q4_2025",
    "Labor+Market+Statistics+Q2+2026": "labour_q2_2026",
}


class RecordedGastat:
    name = "recorded"

    def get(self, url, source_id, params=None, headers=None):
        parts = urlsplit(url)
        if parts.path == "/ar/search":
            query = dict(parse_qsl(parts.query)).get("q", "")
            body = gzip.decompress((SEARCHES / QUERIES[query]).read_bytes()) \
                if query in QUERIES else b""
            return Response(url, url, 200, body, "text/html", {})
        link = BULLETINS.get(unquote(parts.path))
        if link:
            page = f'<a href="https://www.stats.gov.sa{link}">XLSX</a>'.encode()
            return Response(url, url, 200, page, "text/html", {})
        for key, name in WORKBOOKS.items():
            if key in parts.path:
                return Response(url, url, 200, book(name),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml"
                                ".sheet", {})
        return Response(url, url, 200, b"<html></html>", "text/html", {})

    def describe(self):
        return self.name


def ask(question: str, tmp_path: Path):
    http = HttpClient(use_cache=False, retries=1, transport=RecordedGastat())
    registry = Registry(tuple(d for d in load_descriptors() if d.id == "gastat"), http)
    return Agent(registry=registry, http=http, config=AgentConfig(
        out_dir=tmp_path / "out", today=date(2026, 10, 4))).answer(question)


def exported(answer) -> list[list]:
    sheet = load_workbook(answer.primary.export_path)["البيانات"]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


class TestAskingGastat:
    def test_the_unemployment_rate_of_a_year(self, tmp_path):
        answer = ask("معدل البطالة 2025", tmp_path)
        assert answer.verdict is Verdict.AVAILABLE
        primary = answer.primary
        assert primary.candidate.title_ar == "إحصاءات سوق العمل للربع الرابع 2025"
        assert primary.coverage.origin is CoverageOrigin.OBSERVED_DATA
        assert "المعدل الربعي للبطالة" in primary.table_title
        rows = exported(answer)
        assert rows[0][:3] == ["الربع / السنة", "السعوديين - الذكور", "السعوديين - الاناث"]
        assert [r[0] for r in rows[1:]] == ["الربع الأول 2025", "الربع الثاني 2025",
                                            "الربع الثالث 2025", "الربع الرابع 2025", 2025]
        assert "الجدول المستخدم من الملف" in answer.message_ar

    def test_a_count_is_not_answered_with_shares(self, tmp_path):
        # «التوزيع النسبي للمشتغلين (%)» was the answer to «عدد المشتغلين».
        answer = ask("عدد المشتغلين 2025", tmp_path)
        primary = answer.primary
        assert not primary.measure_mismatch
        assert "التوزيع النسبي" not in primary.table_title
        assert primary.candidate.title_ar.startswith("الإحصاءات السجليَّة لسوق العمل")
