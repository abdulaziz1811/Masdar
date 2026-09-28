"""The World Bank's indicators for the Kingdom: an international source.

The data fixture is the API's real answer for electricity use per capita,
recorded on 2026-09-28 (its first four rows, 2025 back to 2022, with the
paging metadata set to one page). 2025 and 2024 come back as `null`: the
bank lists the years but has no figure for them, and they must never be
reported as available.
"""

from __future__ import annotations

import json
from dataclasses import replace
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from masdar.domain.models import Coverage, DatasetCandidate, Verdict
from masdar.nlu.parser import parse
from masdar.pipeline import resolve
from masdar.pipeline.orchestrator import Agent, AgentConfig, _finding_sort_key
from masdar.sources.base import SourceError
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry, load_descriptors
from masdar.sources.transport import Response, Transport

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "worldbank" / "elec_per_capita.json"
ELECTRICITY = "EG.USE.ELEC.KH.PC"


class Recorded(Transport):
    name = "recorded"

    def __init__(self, body: bytes | None = None):
        self.body = body if body is not None else FIXTURE.read_bytes()
        self.urls: list[str] = []

    def get(self, url, source_id, params=None, headers=None):
        self.urls.append(url)
        if f"/indicator/{ELECTRICITY}" in url:
            return Response(url, url, 200, self.body, "application/json", {})
        raise AssertionError(f"unexpected request: {url}")


def worldbank():
    return next(d for d in load_descriptors() if d.id == "worldbank")


def registry(transport=None):
    http = HttpClient(use_cache=False, retries=1, transport=transport or Recorded())
    return Registry((worldbank(),), http)


def payload(rows, **meta):
    head = {"page": 1, "pages": 1, "per_page": 200, "total": len(rows),
            "lastupdated": "2026-07-13", **meta}
    return json.dumps([head, rows], ensure_ascii=False).encode("utf-8")


class TestConfiguration:
    def test_the_source_is_registered_as_international_and_verified(self):
        descriptor = worldbank()
        assert descriptor.international and descriptor.api_verified
        assert descriptor.enabled

    def test_the_bundled_catalogue_loads(self):
        indicators = registry().adapter("worldbank")._indicators()
        assert len(indicators) > 1000
        assert all(i["ar"] and not i["ar"].startswith("اسم المؤشر") for i in indicators
                   if i["ar"])


class TestSearch:
    @pytest.mark.parametrize("question, expected", [
        ("عدد السكان 2024", "SP.POP.TOTL"),
        ("معدل البطالة", "SL.UEM.TOTL.ZS"),
        ("استهلاك الكهرباء للفرد 2023", ELECTRICITY),
        ("الناتج المحلي الاجمالي 2023", "NY.GDP.MKTP.CD"),
        ("عدد السياح", "ST.INT.ARVL"),
        ("معدل الخصوبة", "SP.DYN.TFRT.IN"),
        ("electricity consumption per capita", ELECTRICITY),
    ])
    def test_the_plainest_matching_indicator_comes_first(self, question, expected):
        found = registry().adapter("worldbank").search(parse(question))
        assert found and found[0].dataset_id == expected

    def test_a_ratio_to_gdp_is_not_gdp(self):
        found = registry().adapter("worldbank").search(parse("الناتج المحلي الاجمالي"))
        assert "NE.TRD.GNFS.ZS" not in [c.dataset_id for c in found]  # Trade (% of GDP)

    def test_a_question_the_bank_does_not_cover_finds_nothing(self):
        assert registry().adapter("worldbank").search(parse("كم صيدلية في الرياض")) == []

    def test_candidates_carry_the_source_and_its_caveats(self):
        candidate = registry().adapter("worldbank").search(parse("استهلاك الكهرباء للفرد"))[0]
        assert candidate.publisher_ar == "البنك الدولي"
        assert candidate.license_name == "CC BY 4.0"
        assert 0 < candidate.relevance <= 1
        assert any(c.startswith("مصدر دولي") for c in candidate.caveats)
        assert any("Electric power consumption" in c for c in candidate.caveats)


class TestFetch:
    def _fetch(self, transport=None):
        adapter = registry(transport).adapter("worldbank")
        candidate = adapter.search(parse("استهلاك الكهرباء للفرد"))[0]
        fetched = adapter.fetch(candidate.resources[0].url)
        return candidate, json.loads(fetched.content)

    def test_null_years_are_left_out(self):
        _, rows = self._fetch()
        assert [r["السنة"] for r in rows] == [2022, 2023]
        assert rows[-1]["القيمة"] == pytest.approx(11910.5778104451)

    def test_the_release_date_comes_from_the_response(self):
        candidate, _ = self._fetch(Recorded(payload(
            json.loads(FIXTURE.read_text())[1], lastupdated="2026-09-01")))
        assert candidate.last_updated == date(2026, 9, 1)

    def test_a_series_longer_than_one_page_is_refused(self):
        rows = json.loads(FIXTURE.read_text())[1]
        with pytest.raises(SourceError, match="أطول من صفحة"):
            self._fetch(Recorded(payload(rows, pages=2)))

    def test_an_api_error_is_an_error(self):
        body = json.dumps([{"message": [{"id": "120", "value": "Invalid value"}]}]).encode()
        with pytest.raises(SourceError, match="غير متوقع"):
            self._fetch(Recorded(body))

    def test_a_series_of_nulls_is_not_data(self):
        rows = [dict(r, value=None) for r in json.loads(FIXTURE.read_text())[1]]
        with pytest.raises(SourceError, match="لا توجد قيم"):
            self._fetch(Recorded(payload(rows)))


class TestAnswers:
    def agent(self, tmp_path):
        reg = registry()
        return Agent(registry=reg, http=reg._http,
                     config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 9, 28)))

    def test_a_year_with_a_figure_is_available_and_exported(self, tmp_path):
        answer = self.agent(tmp_path).answer("استهلاك الكهرباء للفرد 2023")
        assert answer.verdict is Verdict.AVAILABLE
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        values = [c for row in sheet.iter_rows(values_only=True) for c in row]
        assert any(isinstance(v, float) and round(v) == 11911 for v in values)

    def test_a_year_listed_as_null_is_not_available(self, tmp_path):
        answer = self.agent(tmp_path).answer("استهلاك الكهرباء للفرد 2024")
        assert answer.verdict is Verdict.NOT_AVAILABLE
        assert 2023 in [s.year for s in answer.suggestions]

    def test_the_answer_says_the_source_is_international(self, tmp_path):
        answer = self.agent(tmp_path).answer("استهلاك الكهرباء للفرد 2023")
        assert any(note.startswith("مصدر دولي") for note in answer.primary.notes)
        assert answer.primary.candidate.last_updated == date(2026, 7, 13)


class TestSaudiFirst:
    def _candidate(self, source_id):
        return DatasetCandidate(source_id=source_id, dataset_id="x", title_ar="عدد السكان")

    def test_an_international_copy_scores_below_the_same_saudi_result(self):
        saudi = replace(worldbank(), id="saudi", international=False)
        request = parse("عدد السكان")
        local = resolve.score_candidate(request, self._candidate("saudi"), saudi)
        foreign = resolve.score_candidate(request, self._candidate("worldbank"), worldbank())
        assert foreign.score < local.score
        assert any("مصدر دولي" in r for r in foreign.match_reasons)

    def _finding(self, source_id, verdict, score):
        candidate = self._candidate(source_id)
        candidate.score = score
        # Only the fields the sort reads.
        return SimpleNamespace(candidate=candidate, verdict=verdict,
                               coverage=Coverage.unknown(), export_path="x.xlsx")

    def test_at_the_same_verdict_the_saudi_finding_comes_first(self):
        foreign = self._finding("worldbank", Verdict.AVAILABLE, 30.0)
        local = self._finding("saudi", Verdict.AVAILABLE, 10.0)
        ordered = sorted([foreign, local], key=lambda f: _finding_sort_key(f, {"worldbank"}))
        assert ordered[0] is local

    def test_a_better_verdict_still_wins(self):
        foreign = self._finding("worldbank", Verdict.AVAILABLE, 10.0)
        local = self._finding("saudi", Verdict.NOT_AVAILABLE, 30.0)
        ordered = sorted([local, foreign], key=lambda f: _finding_sort_key(f, {"worldbank"}))
        assert ordered[0] is foreign


class TestCatalogueBuild:
    def test_only_indicators_with_recent_saudi_values_are_kept(self, tmp_path):
        from masdar.sources.adapters.worldbank import build_catalogue, write_catalogue

        def names(lang):
            label = {"ar": "استهلاك الكهرباء", "en": "Electric power consumption"}[lang]
            return [{"page": 1}, [
                {"id": ELECTRICITY, "name": f"اسم المؤشر{label}" if lang == "ar" else label,
                 "sourceOrganization": "IEA"},
                {"id": "EMPTY.ONE", "name": "مؤشر بلا قيم", "sourceOrganization": ""},
            ]]

        def cell(series, year, value):
            return {"variable": [{"concept": "Series", "id": series},
                                 {"concept": "Time", "id": f"YR{year}"}], "value": value}

        bulk = {"lastupdated": "2026-07-13", "source": {"data": [
            cell(ELECTRICITY, 2023, 11910.5), cell(ELECTRICITY, 2024, None),
            cell("EMPTY.ONE", 2023, None),
        ]}}

        class Api(Transport):
            name = "api"

            def get(self, url, source_id, params=None, headers=None):
                body = (names("ar") if "/ar/source" in url else
                        names("en") if "/en/source" in url else bulk)
                data = json.dumps(body, ensure_ascii=False).encode("utf-8")
                return Response(url, url, 200, data, "application/json", {})

        http = HttpClient(use_cache=False, retries=1, transport=Api())
        catalogue = build_catalogue(http)
        assert [i["id"] for i in catalogue["indicators"]] == [ELECTRICITY]
        entry = catalogue["indicators"][0]
        assert entry["ar"] == "استهلاك الكهرباء" and entry["latest"] == 2023
        assert catalogue["lastupdated"] == "2026-07-13"

        path = tmp_path / "wdi.json"
        write_catalogue(catalogue, path)
        assert json.loads(path.read_text(encoding="utf-8")) == catalogue
