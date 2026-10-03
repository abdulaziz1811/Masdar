"""GASTAT's website search: bulletins with spreadsheets, and data tables.

The search pages under fixtures/gastat_site were recorded on 2026-09-30 from
a German exit (the site answers from outside the Kingdom) and trimmed to the
result list; the bulletin page keeps only its /documents/ links. The
spreadsheet behind the CPI bulletin could not be fetched from the build
environment, so the one served here is SYNTHETIC, built in `_synthetic_cpi`.

The Hajj recordings (2026-10-03, through Firecrawl) are real: the search for
«الحج», the 2026 bulletin's links, and its workbook, rebuilt cell by cell
from GASTAT's file as Firecrawl read it (`Hajj_Statistics_2026_AR.rebuilt.xlsx`:
values verbatim, merged cells and styling not reproduced).
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlsplit

from openpyxl import Workbook

from masdar.domain.models import CoverageOrigin, Verdict
from masdar.nlu.lexicon import load_lexicon
from masdar.nlu.parser import parse
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.adapters.gastat_site import (
    PAGE_NOTE,
    TABLE_NOTE,
    classify,
    document_links,
    parse_results,
    query_words,
)
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry, load_descriptors
from masdar.sources.transport import Response

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "gastat_site"
RECORDED = {
    ("المعتمرين",): "search_umrah.html",
    ("البطالة",): "search_unemployment.html",
    ("التقديرات", "السكانية"): "search_population_estimates.html",
    ("الرقم", "القياسي", "لأسعار", "المستهلك"): "search_cpi.html",
    ("الحجاج",): "search_pilgrims.html",
    ("مؤشرات", "سوق", "العمل"): "search_labour_indicators.html",
    ("الصادرات", "السلعية", "غير", "البترولية"): "search_non_oil_exports.html",
    ("الرقم", "القياسي", "لأسعار", "العقارات"): "search_real_estate_index.html",
    ("المعتمرين", "الخارج"): "search_umrah_abroad.html",
    ("الحج",): "search_hajj.html",
}
CPI_AUGUST = "/ar/w/consumer-price-index-august-2026"
HAJJ_2026 = "/ar/w/إحصاءات-الحج-لعام-2026"
HAJJ_WORKBOOK = FIXTURES / "Hajj_Statistics_2026_AR.rebuilt.xlsx"


def _page(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


def _synthetic_cpi() -> bytes:
    """A stand-in for the bulletin's spreadsheet: formulaic, not GASTAT's."""
    book = Workbook()
    sheet = book.active
    sheet.append(["السنة", "الشهر", "الرقم القياسي العام (بيانات تجريبية)"])
    for month in range(1, 9):
        sheet.append([2026, month, 100 + month])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class SiteTransport:
    """GASTAT's site as recorded: a search answers when its words were recorded."""

    name = "recorded"

    def __init__(self):
        self.urls: list[str] = []

    def get(self, url, source_id, params=None, headers=None):
        self.urls.append(url)
        parts = urlsplit(url)
        if parts.path == "/ar/search":
            asked = set(dict(parse_qsl(parts.query)).get("q", "").split())
            # The recording whose words best cover the question's.
            matching = [(len(words), name) for words, name in RECORDED.items()
                        if set(words) <= asked]
            body = _page(max(matching)[1]) if matching else ""
            return Response(url, url, 200, body.encode(), "text/html", {})
        if unquote(parts.path) == CPI_AUGUST:
            return Response(url, url, 200, _page("bulletin_cpi_august_2026.html").encode(),
                            "text/html", {})
        if unquote(parts.path) == HAJJ_2026:
            return Response(url, url, 200, _page("bulletin_hajj_2026.html").encode(),
                            "text/html", {})
        if "Hajj_Statistics_2026_AR.xlsx" in parts.path:
            return Response(url, url, 200, HAJJ_WORKBOOK.read_bytes(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            {})
        if "/documents/" in parts.path and ".xlsx" in parts.path:
            return Response(url, url, 200, _synthetic_cpi(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            {})
        if parts.path.startswith("/ar/w/"):
            return Response(url, url, 200, b"<html></html>", "text/html", {})
        raise AssertionError(f"unexpected request: {url}")

    def describe(self):
        return self.name


def agent(tmp_path, transport=None):
    http = HttpClient(use_cache=False, retries=1, transport=transport or SiteTransport())
    registry = Registry(tuple(d for d in load_descriptors() if d.id == "gastat"), http)
    return Agent(registry=registry, http=http,
                 config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 9, 30)))


class TestReadingTheSearchPage:
    def test_every_result_is_read(self):
        results = parse_results(_page("search_umrah.html"))
        assert len(results) == 40
        assert all(r["title"] and r["url"].startswith("https://www.stats.gov.sa/") for r in results)
        # The search page's back-link parameters are not part of the link.
        assert not any("p_l_back_url" in r["url"] for r in results)
        assert results[1]["date"] == date(2026, 4, 12)

    def test_tables_and_bulletins_are_told_from_news(self):
        umrah = [classify(r) for r in parse_results(_page("search_umrah.html"))]
        assert umrah.count("table") == 39
        cpi = [classify(r) for r in parse_results(_page("search_cpi.html"))]
        assert cpi.count("bulletin") == 31 and cpi.count("table") == 0
        news = [r for r in parse_results(_page("search_cpi.html"))
                if "/w/news/" in r["url"]]
        assert news and all(classify(r) is None for r in news)

    def test_a_bulletins_spreadsheet_is_found_and_its_pdf_is_not(self):
        files = document_links(_page("bulletin_cpi_august_2026.html"),
                               "https://www.stats.gov.sa" + CPI_AUGUST)
        assert [f.format for f in files] == ["XLSX"]
        assert files[0].title == "CPI Tables-Aug 2026-AR-EN (1).xlsx"
        assert files[0].url.startswith("https://www.stats.gov.sa/documents/20117/")


class TestTheQuery:
    def test_the_askers_words_without_years_or_filler(self):
        words = query_words(parse("كم عدد المعتمرين من الخارج في 2023"), load_lexicon().stopwords)
        assert "المعتمرين" in words and "الخارج" in words
        assert "2023" not in words and "من" not in words and "في" not in words

    def test_spelling_is_kept_for_the_sites_own_search(self):
        words = query_words(parse("الرقم القياسي لأسعار المستهلك"), load_lexicon().stopwords)
        assert words == ["الرقم", "القياسي", "لأسعار", "المستهلك"]


class TestAnswers:
    def test_a_table_seen_on_the_site_is_found_not_denied(self, tmp_path):
        answer = agent(tmp_path).answer("المعتمرين من الخارج حسب منفذ الدخول")
        assert answer.verdict is Verdict.UNVERIFIED
        primary = answer.primary
        assert "منفذ الدخول" in primary.candidate.title_ar
        assert primary.provenance.landing_url.startswith("https://www.stats.gov.sa/ar/w/")
        assert TABLE_NOTE in primary.notes
        assert "لم يُعثر" not in answer.message_ar

    def test_a_bulletins_spreadsheet_is_opened_and_read(self, tmp_path):
        answer = agent(tmp_path).answer("الرقم القياسي لأسعار المستهلك لشهر أغسطس 2026")
        assert answer.verdict is Verdict.AVAILABLE
        primary = answer.primary
        assert "أغسطس 2026" in primary.candidate.title_ar
        assert primary.coverage.origin is CoverageOrigin.OBSERVED_DATA
        assert primary.export_path
        assert primary.provenance.last_updated == date(2026, 9, 15)

    def test_the_unemployment_table_is_found_by_its_subject(self, tmp_path):
        # «معدل» is left out of the site's query: it matches every rate on
        # the site (fertility, productivity...), and the search takes any word.
        answer = agent(tmp_path).answer("معدل البطالة")
        titles = [f.candidate.title_ar for f in answer.findings]
        assert titles and all("البطالة" in t for t in titles)

    def test_the_same_title_twice_is_kept_apart_by_its_survey_year(self, tmp_path):
        answer = agent(tmp_path).answer("المعتمرين من الخارج حسب منفذ الدخول")
        entry = [f for f in answer.findings if "منفذ الدخول" in f.candidate.title_ar]
        years = [f.coverage.years for f in entry]
        # One per survey, told apart by the workbook's name («umrah_2017»).
        assert len(entry) >= 2 and len(set(years)) == len(entry)
        assert all(y and min(y) >= 2015 and max(y) <= 2026 for y in years)
        assert all(f.coverage.origin is CoverageOrigin.INFERRED_TITLE for f in entry)

    def test_umrah_in_a_workbook_name_is_not_a_hijri_marker(self):
        from masdar.nlu.parser import extract_years

        # «ah» inside «umrah» once made 2018 a Hijri year: 2579.
        assert extract_years("umrah survey 2018 0 AR") == (2018,)
        assert extract_years("1445 AH") == (2023, 2024)

    def test_the_site_order_never_stands_in_for_the_subject(self):
        from masdar.pipeline.resolve import SUBJECT_RELEVANCE
        from masdar.sources.adapters.gastat_site import MAX_RELEVANCE

        assert MAX_RELEVANCE < SUBJECT_RELEVANCE

    def test_nothing_found_points_to_gastats_own_search(self, tmp_path):
        answer = agent(tmp_path).answer("كم عدد الخيول")
        assert answer.verdict is Verdict.NO_SOURCE
        (label, url), = answer.elsewhere
        assert "موقع الهيئة العامة للإحصاء" in label
        assert url.startswith("https://www.stats.gov.sa/ar/search?q=")
        assert url in answer.message_ar

    def test_a_bulletin_without_a_file_is_still_a_link(self, tmp_path):
        answer = agent(tmp_path).answer("الرقم القياسي لأسعار المستهلك لشهر مايو 2026")
        may = next(f for f in answer.findings if "مايو 2026" in f.candidate.title_ar)
        assert PAGE_NOTE in may.notes
        assert may.provenance.landing_url.endswith("consumer-price-index-may-2026-1")


class TestWhatTheTesterMet:
    """Questions that once came back «لم يُعثر» though GASTAT publishes the data."""

    def test_an_off_subject_result_does_not_crowd_out_the_answer(self, tmp_path):
        # «المؤشرات الرئيسية» was the only result naming «مؤشرات»; it was kept
        # alone, then dropped as off-subject, and the labour bulletins lost.
        answer = agent(tmp_path).answer("مؤشرات سوق العمل")
        assert answer.verdict is not Verdict.NO_SOURCE
        assert "إحصاءات سوق العمل" in answer.primary.candidate.title_ar

    def test_exports_are_found_under_the_name_gastat_publishes_them(self, tmp_path):
        answer = agent(tmp_path).answer("الصادرات السلعية غير البترولية")
        assert answer.verdict is not Verdict.NO_SOURCE
        assert "التجارة الدولية السلعية غير البترولية" in answer.primary.candidate.title_ar

    def test_the_latest_quarter_comes_first(self, tmp_path):
        answer = agent(tmp_path).answer("الرقم القياسي لأسعار العقارات")
        assert "الربع الثاني 2026" in answer.primary.candidate.title_ar

    def test_the_year_asked_steers_among_surveys(self, tmp_path):
        answer = agent(tmp_path).answer("عدد المعتمرين من الخارج 2018")
        assert 2018 in answer.primary.coverage.years


class TestHajj:
    """«عدد الحجاج»: the site's search knows «الحج», not «الحجاج»."""

    @staticmethod
    def exported(answer):
        from openpyxl import load_workbook

        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    def test_the_bulletin_is_found_by_gastats_word_for_it(self, tmp_path):
        transport = SiteTransport()
        answer = agent(tmp_path, transport).answer("عدد الحجاج 2026")
        searched = [dict(parse_qsl(urlsplit(u).query))["q"] for u in transport.urls
                    if urlsplit(u).path == "/ar/search"]
        assert searched == ["الحجاج", "الحج"]
        assert answer.verdict is Verdict.AVAILABLE
        assert answer.primary.candidate.title_ar == "إحصاءات الحج لعام 2026"

    def test_the_workbooks_link_is_read_as_the_site_now_writes_it(self):
        files = document_links(_page("bulletin_hajj_2026.html"),
                               "https://www.stats.gov.sa" + HAJJ_2026)
        assert [(f.format, f.title) for f in files] == [("XLSX", "Hajj_Statistics_2026_AR.xlsx")]

    def test_the_year_is_read_from_the_tables_title_in_the_file(self, tmp_path):
        answer = agent(tmp_path).answer("عدد الحجاج 2026")
        primary = answer.primary
        assert primary.coverage.origin is CoverageOrigin.OBSERVED_DATA
        assert primary.coverage.years == frozenset({2026})
        assert "عنوان الجدول داخل الملف" in primary.coverage.note

    def test_the_total_is_delivered_not_the_contents_sheet(self, tmp_path):
        rows = self.exported(agent(tmp_path).answer("عدد الحجاج 2026"))
        assert rows[0] == ["جهة القدوم", "الإجمالي"]
        assert ["الإجمالي", 1707301] in rows
        assert not any("رقم الجدول" in str(cell) for row in rows for cell in row)

    def test_the_table_by_sex_answers_the_question_by_sex(self, tmp_path):
        rows = self.exported(agent(tmp_path).answer("عدد الحجاج حسب الجنس 2026"))
        assert [893396, 813905, 1707301] in rows

    def test_one_bulletin_is_no_proof_another_year_is_missing(self, tmp_path):
        # The site holds the 2026 bulletin only; GASTAT published 2023's too.
        answer = agent(tmp_path).answer("عدد الحجاج 2023")
        assert answer.verdict is not Verdict.NOT_AVAILABLE
        assert "غير موجودة" not in answer.message_ar
        assert any("إصدار واحد" in n for n in answer.primary.notes)
        assert 2026 in [s.year for s in answer.suggestions]
