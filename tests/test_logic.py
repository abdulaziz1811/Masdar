"""Logic checks from a battery of real questions (2026-09-29).

Each class is an error the battery found: a result sharing only the topic,
a count answered with a rate, a source crowding out the Saudi publisher, a
region or a follow-up not understood. Searches here run on the local
catalogues; the pipeline tests use the demo sources.
"""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from masdar.chat.server import ChatApp
from masdar.domain.models import DatasetCandidate, Verdict
from masdar.export.tabular import Table
from masdar.nlu.lexicon import load_lexicon
from masdar.nlu.normalize import mentions
from masdar.nlu.parser import parse
from masdar.pipeline import resolve
from masdar.pipeline.orchestrator import Agent, AgentConfig, _each_source_first
from masdar.sources.http import HttpClient
from masdar.sources.registry import DEMO_SOURCES_FILE, Registry, load_descriptors, load_registry

TODAY = date(2026, 9, 27)


@pytest.fixture
def catalogues(monkeypatch):
    """The real registry as the hosted server sees it, searching offline."""
    monkeypatch.setenv("MASDAR_OUTSIDE_KSA", "1")
    http = HttpClient(offline=True, use_cache=False, retries=1)
    return load_registry(http=http)


def ranked_for(registry, question: str) -> list[DatasetCandidate]:
    request = parse(question)
    found = []
    for descriptor in registry.plan(request.topic.id if request.topic else None,
                                    request.topic.preferred_sources if request.topic else ()):
        try:
            found.extend(registry.adapter(descriptor.id).search(request, limit=8))
        except Exception:
            continue  # sources that need the network
    descriptors = {d.id: d for d in registry.descriptors}
    ranked = resolve.rank(request, found, descriptors, load_lexicon(), TODAY)
    return resolve.prefer_specific(request, [c for c in ranked if resolve.on_subject(request, c)])


@pytest.fixture
def demo_agent(tmp_path):
    http = HttpClient(offline=True, use_cache=False, retries=1)
    return Agent(registry=Registry(load_descriptors(DEMO_SOURCES_FILE), http), http=http,
                 config=AgentConfig(out_dir=tmp_path / "out", today=TODAY))


class TestTheConceptAskedMustBeNamed:
    def test_a_topic_alone_is_not_enough(self):
        request = parse("أحدث بيانات البطالة")
        hajj_workforce = DatasetCandidate("x", "1", "القوى العاملة في خدمة الحجاج حسب القطاع")
        unemployment = DatasetCandidate("x", "2", "معدل البطالة حسب الجنس")
        assert not resolve.on_subject(request, hajj_workforce)
        assert resolve.on_subject(request, unemployment)

    def test_word_forms_still_match(self):
        assert mentions("استهلاك الطاقة الكهربائية حسب المنطقة", "استهلاك الكهرباء")
        assert mentions("التقديرات السكانية", "السكان")
        assert mentions("عدد الصيدليات", "صيدلية")
        assert not mentions("الإيرادات التشغيلية للمنشآت", "الصادرات")
        # «ل» for "for", glued to the word: «لمنصة إحسان» names «منصة».
        assert mentions("احصائيات التبرعات لمنصة احسان", "منصة")
        assert mentions("عدد اللقاحات", "لقاحات")

    def test_unemployment_is_found_in_the_world_bank_catalogue(self, catalogues):
        for question in ("أحدث بيانات البطالة", "معدل البطالة حسب الجنس 2023"):
            top = ranked_for(catalogues, question)
            assert top and "بطالة" in top[0].title_ar, question


class TestTheQuantityAsked:
    def test_a_count_of_pilgrims_before_the_workforce_serving_them(self, catalogues):
        top = ranked_for(catalogues, "عدد الحجاج 2023")[0]
        assert top.title_ar.startswith("عدد الحجاج")

    def test_count_tables_survive_the_source_cut(self, catalogues):
        top = ranked_for(catalogues, "عدد المنشآت 2023")[0]
        assert top.title_ar.startswith("عدد المنشآت")

    def test_a_rate_given_for_a_count_is_said(self):
        request = parse("عدد المستشفيات 2022")
        beds = DatasetCandidate("worldbank", "1", "أسِرَّة المستشفيات (لكل 1000 شخص)")
        assert resolve.measure_mismatch(request, beds)


class TestWordsTheDictionaryDoesNotKnow:
    def test_a_broader_result_says_so(self):
        request = parse("نسبة تملك المساكن 2023")
        dwellings = DatasetCandidate("x", "1", "استهلاك الطاقة الكهربائية في المساكن")
        assert resolve.unmatched_terms(request, dwellings) == ("تملك",)

    def test_a_result_naming_the_word_is_preferred(self):
        request = parse("نسبة تملك المساكن 2023")
        dwellings = DatasetCandidate("x", "1", "استهلاك الطاقة الكهربائية في المساكن")
        ownership = DatasetCandidate("x", "2", "نسبة الأسر التي تملك مساكنها")
        assert resolve.prefer_specific(request, [dwellings, ownership]) == [ownership]

    def test_per_capita_goes_to_the_per_capita_indicator(self, catalogues):
        top = ranked_for(catalogues, "استهلاك الكهرباء للفرد 2023")[0]
        assert top.source_id == "worldbank" and "الفرد" in top.title_ar

    def test_publishers_synonyms_count(self):
        request = parse("الصادرات غير النفطية 2023")
        exports = DatasetCandidate("x", "1", "الصادرات السلعية غير البترولية")
        assert resolve.unmatched_terms(request, exports) == ()


class TestEverySourceGetsAFileOpened:
    def test_each_source_first(self):
        wb = [DatasetCandidate("worldbank", str(i), f"مؤشر {i}") for i in range(3)]
        gastat = DatasetCandidate("gastat_cdata", "g", "التقديرات السكانية")
        order = _each_source_first([*wb, gastat])
        assert [c.source_id for c in order[:2]] == ["worldbank", "gastat_cdata"]

    def test_the_saudi_population_table_is_among_the_files_opened(self, catalogues):
        order = _each_source_first(ranked_for(catalogues, "عدد السكان 2022"))
        assert "gastat_cdata" in {c.source_id for c in order[:3]}


class TestRegions:
    def test_a_named_region_asks_for_the_regional_table(self):
        request = parse("استهلاك الكهرباء في الرياض 2022")
        assert request.places == ("الرياض",)
        assert request.free_terms == ()
        assert [d.value for d in request.dimensions] == ["region"]

    def test_the_file_holds_only_that_region(self, demo_agent):
        answer = demo_agent.answer("استهلاك الكهرباء في الرياض 2024")
        assert answer.verdict is Verdict.AVAILABLE
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        rows = list(sheet.iter_rows(values_only=True))
        assert [r[1] for r in rows[1:]] == ["الرياض"]
        assert "اقتصر الملف على: الرياض." in answer.primary.notes

    def test_a_column_that_is_not_about_regions_is_left_alone(self):
        lexicon = load_lexicon()
        stages = Table(["المرحلة", "العدد"],
                       [["رياض الأطفال", 1], ["الابتدائية", 2], ["المتوسطة", 3]])
        table, narrowed = stages.filter_places({"الرياض": lexicon.places["الرياض"]}, lexicon.places)
        assert not narrowed and table is stages


class TestConversation:
    def test_a_region_follow_up(self, demo_agent):
        app = ChatApp.build(demo_agent)
        first = app.ask(None, "استهلاك الكهرباء 2024 حسب المناطق")
        session = first.get("session") or first.get("session_id")
        follow = app.ask(session, "وفي مكة؟")
        assert follow["understood"]["follow_up"] is True
        assert follow["understood"]["places"] == ["مكة المكرمة"]
        assert follow["verdict"] == "available"

    @pytest.mark.parametrize("greeting", ["السلام عليكم", "شكرا", "مرحبا"])
    def test_a_greeting_is_not_searched(self, greeting):
        assert not parse(greeting).is_answerable



class TestWhatGastatsSiteTaught:
    def test_nationality_is_not_sex(self):
        # «حسب الجنس» once chose the table by nationality.
        assert not mentions("أعداد حجاج الداخل حسب الجنسية", "الجنس")
        assert mentions("عدد جنسيات حجاج الخارج", "الجنسية")

    def test_how_often_a_series_is_issued_is_not_its_subject(self):
        from masdar.nlu.normalize import contains_phrase

        assert contains_phrase("الرقم القياسي السنوي للإنتاج الصناعي لعام 2025",
                               "الرقم القياسي للإنتاج الصناعي")

    def test_an_index_base_year_is_not_a_release_year(self):
        # «مؤشر قيمة الصادرات (2000 = 100)» once took the bonus of a release
        # for the year asked, over the exports themselves.
        from masdar.nlu.normalize import period_grain, title_period

        assert period_grain("مؤشر قيمة الصادرات (2000 = 100)") is None
        assert period_grain("الرقم القياسي لأسعار المستهلك سنة الأساس 2018") is None
        assert title_period("الرقم القياسي لأسعار المستهلك أغسطس 2026 (2023=100)") \
            == pytest.approx(2026 + 8 / 12)

    def test_months_and_years_are_not_sent_to_the_sites_search(self):
        from masdar.sources.adapters.gastat_site import query_words

        words = query_words(parse("معدل التضخم لشهر أغسطس 2026"), load_lexicon().stopwords)
        assert "أغسطس" not in words and "2026" not in words and "التضخم" in words

    def test_gastats_own_name_for_the_subject(self):
        assert resolve.subject_aliases(parse("كم عدد الحجاج")) == ("الحج",)
        assert resolve.subject_aliases(parse("معدل البطالة 2025")) == ("سوق العمل",)
        assert resolve.subject_aliases(parse("نسبة التضخم")) == ("أسعار المستهلك",)

    def test_the_headline_names_what_was_asked_not_its_topic(self):
        from masdar.pipeline.reply import _subject

        assert _subject(parse("عدد الحجاج 2026")) == "عدد الحجاج"
        assert _subject(parse("عطني معدل البطالة لسنة 2025")) == "معدل البطالة"


class TestNotAvailableNeedsEveryLead:
    """«غير متوفرة» is a claim about the data, not about one dataset."""

    @staticmethod
    def finding(source_id, title, verdict, years=()):
        from datetime import datetime

        from masdar.domain.models import Coverage, CoverageOrigin, Finding, Provenance

        coverage = (Coverage(frozenset(years), CoverageOrigin.OBSERVED_DATA, True)
                    if years else Coverage.unknown())
        return Finding(
            candidate=DatasetCandidate(source_id, title, title),
            provenance=Provenance(source_id=source_id, publisher_ar="", publisher_en="",
                                  landing_url="", retrieved_at=datetime(2026, 10, 3)),
            verdict=verdict, coverage=coverage,
        )

    def held(self, question, today=date(2026, 10, 3)):
        from masdar.pipeline.orchestrator import _absence_needs_every_lead

        series = self.finding("worldbank", "Unemployment, total", Verdict.NOT_AVAILABLE,
                              range(2010, 2025))
        table = self.finding("gastat", "معدل البطالة حسب الجنس", Verdict.UNVERIFIED)
        return _absence_needs_every_lead([series, table], parse(question), today,
                                         frozenset({"worldbank"}))

    def test_a_saudi_lead_left_open_keeps_the_answer_unconfirmed(self):
        # The World Bank's series ends at 2024; GASTAT's table is not dated.
        assert self.held("معدل البطالة 2025")[0].candidate.source_id == "gastat"

    def test_a_year_not_yet_over_is_still_not_available(self):
        assert self.held("معدل البطالة 2026")[0].verdict is Verdict.NOT_AVAILABLE


class TestWhatTheOctoberBatteryFound:
    """Faults found by a 45-question battery on 2026-10-17, one test each."""

    def test_a_percent_share_is_no_count(self):
        # normalize() turns «%» into a space, so the sign was never seen.
        request = parse("عدد السكان 2024")
        share = DatasetCandidate("x", "1", "تعداد السكان، إناث (% من الإجمالي)")
        assert resolve.measure_mismatch(request, share)
        assert resolve.measure_mismatch(request, DatasetCandidate("x", "2", "الكثافة السكانية"))
        assert not resolve.measure_mismatch(request, DatasetCandidate("x", "3", "عدد السكان"))

    def test_a_share_unasked_ranks_below_the_amount(self):
        request = parse("إنتاج الكهرباء 2024")
        amount = DatasetCandidate("x", "1", "إنتاج الكهرباء")
        share = DatasetCandidate("x", "2", "إنتاج الكهرباء من المصادر النفطية (% من الإجمالي)")
        scored = {c.dataset_id: resolve.score_candidate(request, c, today=TODAY) for c in
                  (amount, share)}
        assert scored["2"].score < scored["1"].score
        assert any("نسبة مئوية" in r for r in scored["2"].match_reasons)
        # Asked for a share, it is not marked down.
        asked = resolve.score_candidate(parse("نسبة إنتاج الكهرباء من النفط"), share, today=TODAY)
        assert not any("نسبة مئوية" in r for r in asked.match_reasons)

    def test_a_fused_lam_still_names_the_topic(self):
        request = parse("الرقم القياسي لأسعار المستهلك أغسطس 2026")
        assert request.topic.id == "prices"
        assert "أسعار المستهلك" in request.typed_phrases
        # The month is when, not what; «لأسعار» is spent with its topic.
        assert request.free_terms == ()

    def test_the_month_asked_ranks_its_release_first(self):
        request = parse("الرقم القياسي لأسعار المستهلك لشهر مايو 2026")
        may = DatasetCandidate("x", "5", "الرقم القياسي لأسعار المستهلك لشهر مايو 2026")
        august = DatasetCandidate("x", "8", "الرقم القياسي لأسعار المستهلك لشهر أغسطس 2026")
        ranked = resolve.rank(request, [august, may], {}, load_lexicon(), TODAY)
        assert ranked[0].dataset_id == "5"
        assert resolve.other_month(request, august) == frozenset({8})
        assert resolve.other_month(request, may) == frozenset()
        assert resolve.other_month(parse("الرقم القياسي لأسعار المستهلك 2026"), august) \
            == frozenset()

    def test_another_subject_in_an_indicator_name_costs_more(self):
        # «عدد وفيات المواليد» (infant deaths) came first for «عدد المواليد».
        from masdar.sources.adapters.worldbank import _changes_subject

        assert _changes_subject("فيات", ["عدد", "مواليد"])
        # «الطاقة» in «الطاقة الكهربائية» is half of electricity's own name.
        assert not _changes_subject("طاقه", ["استهلاك", "كهرباء"])

    @pytest.mark.parametrize("question, flow", [
        ("عدد الأطباء 2023", "MOBILE_DF_HTH_PHYSICIANS"),
        ("عدد المدارس", "MOBILE_DF_EDU_SCHOOL"),
        ("حالات الطلاق 2021", "MOBILE_DF_DIV_DIVORCE_CASES"),
        ("عدد عقود الزواج 2021", "MOBILE_DF_MAR_MARRIAGE_CASES"),
        ("عدد زوار المتاحف", "MOBILE_DF_CUL_MUES_VIST"),
    ])
    def test_questions_with_no_answer_offline_now_reach_gcc_stat(self, catalogues, question,
                                                                 flow):
        top = ranked_for(catalogues, question)
        assert flow in [c.dataset_id for c in top[:3]]
