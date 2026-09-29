"""End-to-end tests over the fixture sources.

These run with `offline=True`, so a test that passes has provably not
touched the network: everything comes from the fixtures via `file://`.
"""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from masdar.domain.models import Verdict
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.base import SourceDescriptor
from masdar.sources.http import HttpClient
from masdar.sources.registry import DEMO_SOURCES_FILE, Registry, load_descriptors

TODAY = date(2026, 9, 27)


@pytest.fixture
def agent(tmp_path):
    http = HttpClient(offline=True, use_cache=False, retries=1)
    registry = Registry(load_descriptors(DEMO_SOURCES_FILE), http)
    config = AgentConfig(out_dir=tmp_path / "out", today=TODAY)
    return Agent(registry=registry, http=http, config=config)


class TestUnavailableYear:
    """The motivating case: data for a year that is not published."""

    def test_reports_not_available_without_guessing(self, agent):
        answer = agent.answer("ابي احصاءات الطاقه الكهربائيه لسنه 2026 حسب المناطق")
        assert answer.verdict is Verdict.NOT_AVAILABLE
        assert "لا تتوفر بيانات" in answer.message_ar

    def test_verdict_rests_on_the_file_not_the_title(self, agent):
        answer = agent.answer("الكهرباء 2026 حسب المناطق")
        assert answer.primary.coverage.origin.is_authoritative

    def test_offers_the_nearest_available_year(self, agent):
        answer = agent.answer("الكهرباء 2026 حسب المناطق")
        assert [s.year for s in answer.suggestions] == [2024]

    def test_explains_why_the_year_is_missing(self, agent):
        answer = agent.answer("الكهرباء 2026 حسب المناطق")
        assert any("لم تنتهِ بعد" in note for note in answer.primary.notes)

    def test_writes_no_workbook_for_a_year_it_does_not_have(self, agent):
        answer = agent.answer("الكهرباء 2026 حسب المناطق")
        assert answer.primary.export_path is None


class TestAvailableYear:
    def test_confirms_and_exports(self, agent):
        answer = agent.answer("احصاءات الطاقه الكهربائيه لسنه 2024 حسب المناطق")
        assert answer.verdict is Verdict.AVAILABLE
        assert answer.primary.export_path is not None
        assert Path(answer.primary.export_path).exists()

    def test_provenance_is_complete(self, agent):
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        provenance = answer.primary.provenance
        assert provenance.sha256 and len(provenance.sha256) == 64
        assert provenance.landing_url
        assert provenance.resource_url
        assert provenance.last_updated == date(2025, 6, 30)
        assert provenance.retrieved_at is not None

    def test_workbook_has_data_and_source_sheets(self, agent):
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        workbook = load_workbook(answer.primary.export_path)
        assert workbook.sheetnames == ["البيانات", "المصدر"]

    def test_workbook_holds_only_the_requested_year(self, agent):
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        years = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert years == {2024}

    def test_numbers_are_exported_as_numbers(self, agent):
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        first = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        assert isinstance(first[2], int)

    def test_source_sheet_records_provenance(self, agent):
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["المصدر"]
        text = "\n".join(
            str(cell) for row in sheet.iter_rows(values_only=True) for cell in row if cell
        )
        assert answer.primary.provenance.sha256 in text
        assert "2025-06-30" in text
        assert "آخر تحديث معلن من المصدر" in text

    def test_synthetic_data_is_labelled_in_the_workbook(self, agent):
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["المصدر"]
        text = "\n".join(
            str(cell) for row in sheet.iter_rows(values_only=True) for cell in row if cell
        )
        assert "ليست بيانات رسمية" in text


class TestPartialRange:
    def test_reports_both_sides_of_the_gap(self, agent):
        answer = agent.answer("الكهرباء من 2023 الى 2026 حسب المناطق")
        assert answer.verdict is Verdict.PARTIAL
        assert answer.primary.matched_years == (2023, 2024)
        assert answer.primary.missing_years == (2025, 2026)

    def test_workbook_is_named_for_what_it_contains(self, agent):
        answer = agent.answer("الكهرباء من 2023 الى 2026 حسب المناطق")
        assert "2023-2024" in Path(answer.primary.export_path).name

    def test_workbook_excludes_the_missing_years(self, agent):
        answer = agent.answer("الكهرباء من 2023 الى 2026 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        years = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert years == {2023, 2024}


class TestMissingBreakdown:
    def test_warns_when_the_requested_split_is_absent(self, agent):
        # The electricity fixture has no gender column.
        answer = agent.answer("الكهرباء 2024 حسب الجنس")
        assert any("لا يحتوي عموداً للتفصيل المطلوب" in n for n in answer.primary.notes)


class TestUnreachableSources:
    """Being blocked must never be reported as an absence of data."""

    def test_reports_unreachable_not_missing(self, tmp_path):
        http = HttpClient(offline=True, use_cache=False, retries=1)
        blocked = SourceDescriptor(
            id="blocked",
            name_ar="مصدر محجوب",
            name_en="Blocked source",
            base_url="https://blocked.invalid",
            adapter="html_index",
            topics=("electricity",),
        )
        agent = Agent(
            registry=Registry((blocked,), http),
            http=http,
            config=AgentConfig(out_dir=tmp_path / "out", today=TODAY),
        )
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        assert answer.verdict is Verdict.SOURCE_UNREACHABLE
        assert answer.source_errors
        assert "لا يعني أن البيانات غير موجودة" in answer.message_ar

    def test_a_working_source_still_answers_when_another_fails(self, tmp_path):
        http = HttpClient(offline=True, use_cache=False, retries=1)
        blocked = SourceDescriptor(
            id="blocked",
            name_ar="مصدر محجوب",
            name_en="Blocked source",
            base_url="https://blocked.invalid",
            adapter="html_index",
            authority=100,
            topics=("electricity",),
        )
        descriptors = (blocked,) + load_descriptors(DEMO_SOURCES_FILE)
        agent = Agent(
            registry=Registry(descriptors, http),
            http=http,
            config=AgentConfig(out_dir=tmp_path / "out", today=TODAY),
        )
        answer = agent.answer("الكهرباء 2024 حسب المناطق")
        assert answer.verdict is Verdict.AVAILABLE
        # The outage is still reported rather than hidden by the success.
        assert any(source == "blocked" for source, _ in answer.source_errors)


class TestUnparseableRequest:
    def test_asks_for_clarification_instead_of_searching(self, agent):
        answer = agent.answer("ابي شي")
        assert answer.verdict is Verdict.NO_SOURCE
        assert "لم يتضمن الطلب موضوعاً" in answer.message_ar


class TestSourceIsolation:
    def test_demo_sources_are_not_in_the_default_registry(self):
        ids = {d.id for d in load_descriptors()}
        assert not any(i.startswith("demo_") for i in ids)

    def test_no_default_source_is_marked_synthetic(self):
        assert all(not d.synthetic for d in load_descriptors())
