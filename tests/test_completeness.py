"""The exported numbers are the published numbers: complete and unaggregated.

Verified against the live API on 2026-09-28: a request that leaves a
dimension out gets it SUMMED away. For life expectancy grouped by YEAR alone,
2022 comes back as 234.09 -- female 80.89 + total 77.90 + male 75.30.
Both fixtures here are those responses, captured verbatim.
"""

import json
from datetime import date
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlsplit

import pytest
from openpyxl import load_workbook

from masdar.domain.models import Verdict
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.base import SourceError
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry
from masdar.sources.transport import Response
from tests.fakes import aggregate, cdata_descriptor, paged

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "gastat_cdata"
SUMMED_2022 = "234.09"


def dimensions_in(url: str) -> list[str]:
    return [v for k, v in parse_qsl(urlsplit(url).query) if k == "dimensions[]"]


class HealthTransport:
    """Behaves like the live API: year-only requests get the server's sum."""

    name = "recorded"

    def __init__(self, server_cap=None, ignore_skip=False):
        self.urls: list[str] = []
        self.server_cap = server_cap
        self.ignore_skip = ignore_skip

    def get(self, url, source_id, params=None, headers=None):
        self.urls.append(url)
        # Serve what the server would: the published rows, summed over every
        # dimension the request left out.
        full = json.loads((FIXTURES / "HLTH1304_full.json").read_text(encoding="utf-8"))
        rows = aggregate(full["value"], dimensions_in(url))
        raw = json.dumps({"value": rows}, ensure_ascii=False).encode("utf-8")
        page_url = url.replace("$skip=", "$ignored=") if self.ignore_skip else url
        body = paged(raw, page_url, self.server_cap)
        return Response(url, url, 200, body, "application/json", {})

    def describe(self):
        return self.name


def build(tmp_path, transport):
    descriptors = (cdata_descriptor(),)
    http = HttpClient(use_cache=False, retries=1, transport=transport)
    return Agent(
        registry=Registry(descriptors, http),
        http=http,
        config=AgentConfig(out_dir=tmp_path / "out", today=date(2026, 9, 28)),
    )


def workbook_text(path) -> str:
    book = load_workbook(path)
    return "\n".join(
        str(c) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
        for c in row if c is not None
    )


QUERY = "ابي متوسط العمر المتوقع للسعوديين لسنه 2022 حسب الجنس"


class TestTheSimulatorMatchesTheServer:
    def test_reproduces_the_real_sum_exactly(self):
        # The fake is only trustworthy if it reproduces the live response.
        full = json.loads((FIXTURES / "HLTH1304_full.json").read_text(encoding="utf-8"))
        live = json.loads((FIXTURES / "HLTH1304_years_summed.json").read_text(encoding="utf-8"))
        summed = aggregate(full["value"], ["YEAR"])
        simulated = {r["YEAR_TIME"]: r["TOTAL_OBS_VALUE"] for r in summed}
        for row in live["value"]:
            if row["YEAR_TIME"] in simulated:
                assert simulated[row["YEAR_TIME"]] == pytest.approx(row["TOTAL_OBS_VALUE"])


class TestNoServerSideAggregation:
    def test_every_dimension_is_requested_whatever_was_asked(self, tmp_path):
        transport = HealthTransport()
        answer = build(tmp_path, transport).answer(QUERY)
        data_urls = [u for u in transport.urls if dimensions_in(u) != ["YEAR"]]
        assert data_urls, "the data itself was never requested"
        for url in data_urls:
            assert set(dimensions_in(url)) == {"INDICATORS", "SEX", "YEAR"}
        assert answer.verdict is Verdict.AVAILABLE

    def test_published_values_reach_the_workbook(self, tmp_path):
        answer = build(tmp_path, HealthTransport()).answer(QUERY)
        text = workbook_text(answer.primary.export_path)
        for published in ("80.888", "77.904", "75.300"):
            assert published in text

    def test_a_question_with_no_breakdown_still_gets_published_values(self, tmp_path):
        # Asking for no split at all is where a partial request would have
        # summed female + total + male into a 234-year life expectancy.
        answer = build(tmp_path, HealthTransport()).answer(
            "ابي متوسط العمر المتوقع للسعوديين لسنه 2022"
        )
        text = workbook_text(answer.primary.export_path)
        assert SUMMED_2022 not in text
        assert "77.904" in text

    def test_the_servers_sum_never_reaches_the_workbook(self, tmp_path):
        # 234.09 "years" is what an aggregated request returns. It was seen
        # by the coverage probe, which only reads which years exist.
        answer = build(tmp_path, HealthTransport()).answer(QUERY)
        assert SUMMED_2022 not in workbook_text(answer.primary.export_path)

    def test_connections_dataset_url_carries_all_five_dimensions(self, tmp_path):
        from masdar.nlu.parser import parse
        from tests.test_gastat_cdata import RoutingTransport

        transport = RoutingTransport()
        http = HttpClient(use_cache=False, retries=1, transport=transport)
        descriptors = (cdata_descriptor(),)
        found = Registry(descriptors, http).adapter("gastat_cdata").search(
            parse("المساكن المتصلة بالكهرباء 2022 حسب المناطق")
        )
        connections = next(c for c in found if c.dataset_id == "DPV_HES_EHE_IT_HES0301")
        # Asked only for regions; the other four would otherwise be summed.
        assert set(dimensions_in(unquote(connections.resources[0].url))) == {
            "ELECTRICITY_NETWORK_TYPE", "HOUSE_TYPE", "POSSESSION_TYPE", "REGION", "YEAR",
        }


class TestPaging:
    @pytest.fixture(autouse=True)
    def small_pages(self, monkeypatch):
        import masdar.sources.adapters.gastat_cdata as cdata

        monkeypatch.setattr(cdata, "PAGE_SIZE", 2)

    def test_all_rows_arrive_across_pages(self, tmp_path):
        transport = HealthTransport()
        answer = build(tmp_path, transport).answer(QUERY)
        pages = [u for u in transport.urls if "$skip=" in u and dimensions_in(u) != ["YEAR"]]
        assert len(pages) >= 3  # 6 rows at 2 a page, then the empty page
        assert "80.888" in workbook_text(answer.primary.export_path)

    def test_a_server_cap_below_the_page_size_loses_nothing(self, tmp_path, monkeypatch):
        import masdar.sources.adapters.gastat_cdata as cdata

        monkeypatch.setattr(cdata, "PAGE_SIZE", 5)
        # The server silently returns at most 2 rows whatever $top says.
        answer = build(tmp_path, HealthTransport(server_cap=2)).answer(
            "متوسط العمر المتوقع للسعوديين لسنه 2021 حسب الجنس"
        )
        text = workbook_text(answer.primary.export_path)
        for published in ("79.163", "76.007", "73.303"):
            assert published in text

    def test_unstable_pages_are_refused_not_exported(self, tmp_path):
        from masdar.nlu.parser import parse

        # A server that ignores $skip returns the same page forever.
        transport = HealthTransport(ignore_skip=True)
        descriptors = (cdata_descriptor(),)
        http = HttpClient(use_cache=False, retries=1, transport=transport)
        adapter = Registry(descriptors, http).adapter("gastat_cdata")
        candidate = next(
            c for c in adapter.search(parse(QUERY)) if c.dataset_id == "DPV_HLTH13_DPHLTH1304"
        )
        with pytest.raises(SourceError) as caught:
            adapter.fetch(candidate.resources[0].url)
        assert "مكرر" in caught.value.reason

    def test_too_many_pages_is_an_error_not_a_truncation(self, tmp_path, monkeypatch):
        import masdar.sources.adapters.gastat_cdata as cdata
        from masdar.nlu.parser import parse

        monkeypatch.setattr(cdata, "MAX_PAGES", 2)
        descriptors = (cdata_descriptor(),)
        http = HttpClient(use_cache=False, retries=1, transport=HealthTransport())
        adapter = Registry(descriptors, http).adapter("gastat_cdata")
        candidate = next(
            c for c in adapter.search(parse(QUERY)) if c.dataset_id == "DPV_HLTH13_DPHLTH1304"
        )
        with pytest.raises(SourceError) as caught:
            adapter.fetch(candidate.resources[0].url)
        assert "رُفض التصدير" in caught.value.reason

    def test_provenance_cites_the_canonical_query_and_hashes_every_row(self, tmp_path):
        import hashlib

        answer = build(tmp_path, HealthTransport()).answer(QUERY)
        provenance = answer.primary.provenance
        assert "$skip" not in provenance.resource_url
        assert "$top" not in provenance.resource_url
        assert len(provenance.sha256) == 64
        assert provenance.sha256 != hashlib.sha256(b"").hexdigest()


class TestHealthSpecFromThePortal:
    def test_the_five_health_datasets_are_known(self):
        from masdar.sources.http import HttpClient

        descriptors = (cdata_descriptor(),)
        adapter = Registry(descriptors, HttpClient(offline=True, retries=1)).adapter("gastat_cdata")
        ids = {d["id"] for d in adapter._datasets()}
        assert {
            "DPV_HLTH13_DPHLTH130132", "DPV_HLTH13_DPHLTH1304", "DPV_HLTH13_DPHLTH130789",
            "DPV_HLTH13_DPHLTH1311", "DPV_HLTH13_DPHLTH1314",
        } <= ids

    def test_obs_value_measures_are_recognised(self):
        from masdar.sources.http import HttpClient

        descriptors = (cdata_descriptor(),)
        adapter = Registry(descriptors, HttpClient(offline=True, retries=1)).adapter("gastat_cdata")
        life = next(d for d in adapter._datasets() if d["id"] == "DPV_HLTH13_DPHLTH1304")
        assert life["measures"] == ["TOTAL_OBS_VALUE"]

    def test_life_expectancy_question_finds_its_dataset(self, tmp_path):
        answer = build(tmp_path, HealthTransport()).answer(QUERY)
        assert answer.primary.candidate.dataset_id == "DPV_HLTH13_DPHLTH1304"
