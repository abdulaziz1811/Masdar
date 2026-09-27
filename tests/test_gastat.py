"""Tests for the GASTAT API adapter and the JSON table reader.

The fixtures mirror the shapes verified against the live service on
2026-09-27: a bare-array indicator catalog, an OData observations payload
whose columns are named per indicator (`*_TIME`, `*_OBSV`), and the
`{"info": [...]}` headline response captured directly from the API.
"""

from datetime import date
from pathlib import Path

import pytest

from masdar.domain.models import Verdict
from masdar.export.tabular import YearAxis, read_json
from masdar.pipeline.orchestrator import Agent, AgentConfig
from masdar.sources.base import SourceDescriptor
from masdar.sources.http import HttpClient
from masdar.sources.registry import Registry

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "gastat"
TODAY = date(2026, 9, 27)


def descriptor() -> SourceDescriptor:
    """A GASTAT-shaped source served from local files.

    file:// URLs keep a query string in `query`, not `path`, so the
    `?api=<token>` the adapter appends resolves to the same fixture -- the
    request travels the real code path.
    """
    return SourceDescriptor(
        id="gastat_db",
        name_ar="قاعدة بيانات الهيئة العامة للإحصاء",
        name_en="GASTAT statistical database",
        base_url=FIXTURES.as_uri(),
        adapter="gastat_api",
        authority=100,
        api_verified=True,
        api={
            "catalog": "/catalog.json",
            "headlines": "/cardsinfo.json",
            "observations": "/observations.json",
        },
        topics=("electricity", "energy", "population"),
    )


@pytest.fixture
def agent(tmp_path):
    http = HttpClient(offline=True, use_cache=False, retries=1)
    registry = Registry((descriptor(),), http)
    return Agent(
        registry=registry,
        http=http,
        config=AgentConfig(out_dir=tmp_path / "out", today=TODAY),
    )


class TestJsonReader:
    def test_reads_odata_value_array(self):
        table = read_json((FIXTURES / "observations.json").read_bytes())
        assert len(table) == 6
        assert "ELEC_TIME" in table.columns

    def test_drops_odata_bookkeeping_keys(self):
        table = read_json((FIXTURES / "observations.json").read_bytes())
        assert not any(c.startswith("@odata") for c in table.columns)

    def test_period_column_comes_first(self):
        table = read_json((FIXTURES / "observations.json").read_bytes())
        assert table.columns[0] == "ELEC_TIME"

    def test_finds_years_in_a_per_indicator_column_name(self):
        # The column is ELEC_TIME, not "year": detection must go by content.
        table = read_json((FIXTURES / "observations.json").read_bytes())
        assert table.year_axis() is YearAxis.ROWS
        assert table.observed_years() == frozenset({2022, 2023, 2024})

    def test_values_become_numbers(self):
        table = read_json((FIXTURES / "observations.json").read_bytes())
        assert table.rows[0][1] == 1200.5

    def test_reads_the_headline_shape_captured_live(self):
        table = read_json((FIXTURES / "cardsinfo.json").read_bytes())
        assert "ar_title" in table.columns
        assert table.rows[0][table.columns.index("obsvalue")] == 35_013_414

    def test_bare_array_is_accepted(self):
        assert len(read_json(b'[{"a": 1}, {"a": 2}]')) == 2

    def test_empty_payload_is_not_an_error(self):
        assert len(read_json(b'{"value": []}')) == 0


class TestCatalogSearch:
    def _adapter(self):
        http = HttpClient(offline=True, use_cache=False, retries=1)
        return Registry((descriptor(),), http).adapter("gastat_db")

    def test_finds_the_matching_indicator(self):
        from masdar.nlu.parser import parse

        found = self._adapter().search(parse("الكهرباء 2024 حسب المناطق"))
        assert [c.dataset_id for c in found] == ["101"]

    def test_skips_branch_nodes_without_a_data_token(self):
        from masdar.nlu.parser import parse

        # Entry 102 is a folder: it matches on title but carries api_url: null.
        found = self._adapter().search(parse("إحصاءات الطاقة 2024"))
        assert all(c.dataset_id != "102" for c in found)

    def test_builds_the_observations_url_with_the_token(self):
        from masdar.nlu.parser import parse

        candidate = self._adapter().search(parse("الكهرباء 2024"))[0]
        assert "api=TOKEN-ELEC-101" in candidate.resources[0].url

    def test_claims_no_coverage_from_periodicity(self):
        from masdar.nlu.parser import parse

        # frequency_en is "Annual" -- that is a cadence, not a list of years.
        candidate = self._adapter().search(parse("الكهرباء 2024"))[0]
        assert candidate.claimed_coverage.is_empty

    def test_unrelated_query_returns_nothing(self):
        from masdar.nlu.parser import parse

        assert self._adapter().search(parse("أسعار العقار 2024")) == []

    def test_probe_reports_the_headline_endpoint(self):
        assert "Population" in self._adapter().probe()


class TestEndToEndOverGastatShapes:
    def test_available_year_is_confirmed_from_the_observations(self, agent):
        answer = agent.answer("ابي استهلاك الكهرباء لسنه 2024 حسب المناطق")
        assert answer.verdict is Verdict.AVAILABLE
        assert answer.primary.coverage.origin.is_authoritative
        assert answer.primary.matched_years == (2024,)

    def test_exports_a_cited_workbook(self, agent):
        answer = agent.answer("استهلاك الكهرباء لسنه 2024 حسب المناطق")
        assert answer.primary.export_path
        assert Path(answer.primary.export_path).exists()
        assert len(answer.primary.provenance.sha256) == 64

    def test_workbook_holds_only_the_requested_year(self, agent):
        from openpyxl import load_workbook

        answer = agent.answer("استهلاك الكهرباء لسنه 2024 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["البيانات"]
        years = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert years == {2024}

    def test_unpublished_year_is_refused_with_an_alternative(self, agent):
        answer = agent.answer("ابي استهلاك الكهرباء لسنه 2026 حسب المناطق")
        assert answer.verdict is Verdict.NOT_AVAILABLE
        assert [s.year for s in answer.suggestions] == [2024]
        assert answer.primary.export_path is None

    def test_not_labelled_as_sample_data(self, agent):
        from openpyxl import load_workbook

        # This source is real, so the synthetic warning must be absent.
        answer = agent.answer("استهلاك الكهرباء لسنه 2024 حسب المناطق")
        sheet = load_workbook(answer.primary.export_path)["المصدر"]
        text = "\n".join(
            str(c) for row in sheet.iter_rows(values_only=True) for c in row if c
        )
        assert "ليست بيانات رسمية" not in text


class TestCredentialHygiene:
    def test_registry_names_where_secrets_live_not_their_values(self):
        from masdar.sources.registry import load_descriptors

        cdata = next(d for d in load_descriptors() if d.id == "gastat_cdata")
        assert cdata.auth.get("key_env") == "GASTAT_API_KEY"
        # Variable names and a header name only -- nothing credential-shaped.
        for value in cdata.auth.values():
            assert len(str(value)) < 40

    def test_no_key_is_sent_when_the_environment_has_none(self, monkeypatch):
        from masdar.sources.http import HttpClient
        from masdar.sources.registry import Registry, load_descriptors

        monkeypatch.delenv("GASTAT_API_KEY", raising=False)
        descriptors = tuple(d for d in load_descriptors() if d.id == "gastat_cdata")
        adapter = Registry(descriptors, HttpClient(offline=True, retries=1)).adapter(
            "gastat_cdata"
        )
        assert adapter._auth_headers() == {}

    def test_a_configured_key_is_read_from_the_environment(self, monkeypatch):
        from masdar.sources.http import HttpClient
        from masdar.sources.registry import Registry, load_descriptors

        monkeypatch.setenv("GASTAT_API_KEY", "test-value-not-a-real-key")
        descriptors = tuple(d for d in load_descriptors() if d.id == "gastat_cdata")
        adapter = Registry(descriptors, HttpClient(offline=True, retries=1)).adapter(
            "gastat_cdata"
        )
        assert adapter._auth_headers() == {"apikey": "test-value-not-a-real-key"}

    def test_config_contains_no_long_opaque_tokens(self):
        """Guards against a key being pasted into the repo by accident."""
        import re

        from masdar.nlu.lexicon import CONFIG_DIR

        suspicious = re.compile(r"[A-Za-z0-9]{40,}")
        for path in CONFIG_DIR.glob("*.yaml"):
            for match in suspicious.findall(path.read_text(encoding="utf-8")):
                raise AssertionError(f"possible secret in {path.name}: {match[:12]}...")
