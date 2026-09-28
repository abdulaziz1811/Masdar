"""A first look at an exported file, for the chat page.

The preview is read back from the workbook the user will download, not from
the table in memory, so what the page shows is exactly what the file holds.

Charts are drawn only where the file leaves no choice to make:

* **Over the years**: a year column, and every other non-numeric column
  holding a single value (one slice, not a breakdown). One chart per
  numeric column.
* **Across a breakdown, for one year**: the latest year's rows, one
  breakdown (its Arabic, English and code columns may move together), and
  one chart per numeric column. Rows that are the national total
  ("الإجمالي", "المملكة") are left out, and the chart says so, because a
  total drawn beside its parts dwarfs them.

Anything else -- two breakdowns at once, a year repeated within one slice --
gets the table preview alone: a chart that quietly summed or picked a slice
could mislead.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

from masdar.export.excel import DATA_SHEET
from masdar.nlu.normalize import normalize

PREVIEW_ROWS = 8
MAX_CELL_CHARS = 80
# Files longer than this are counted but not scanned for a chart.
MAX_SCAN_ROWS = 5000
MAX_POINTS = 60
MAX_CATEGORIES = 25
MAX_CHARTS = 3
_YEAR_WORDS = ("سنه", "السنه", "العام", "عام", "year", "yr", "time_period", "year_time")
_MEASURE_WORDS = ("القيمه", "قيمه", "value", "obs_value", "obsvalue")
_TOTAL_WORDS = frozenset(normalize(w) for w in (
    "الإجمالي", "إجمالي", "المجموع", "مجموع", "الكل", "المملكة", "إجمالي المملكة",
    "المملكة العربية السعودية", "total", "all", "saudi arabia", "kingdom", "ksa",
))
_ARABIC = re.compile("[\u0600-\u06ff]")
_API_MEASURE_SUFFIXES = ("_OBS_VALUE", "_OBSV", "OBS_VALUE", "OBSVALUE")


def _cell(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value)
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def _year(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int) and 1900 <= value <= 2100:
        return value
    if isinstance(value, float) and value.is_integer() and 1900 <= value <= 2100:
        return int(value)
    if isinstance(value, str) and value.strip().isdigit() and len(value.strip()) == 4:
        return _year(int(value.strip()))
    return None


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _year_column(columns: list[str]) -> int | None:
    for index, name in enumerate(columns):
        # GASTAT's API names it <DIMENSION>_TIME (YEAR_TIME, POP_TIME).
        if normalize(name) in _YEAR_WORDS or name.strip().upper().endswith("_TIME"):
            return index
    return None


def _is_code(name: str) -> bool:
    raw = name.strip().upper()
    return raw.endswith("_CODE") or raw in ("CODE", "ID") or "رمز" in normalize(name)


def _is_measure_name(name: str) -> bool:
    raw = name.strip().upper()
    return normalize(name) in _MEASURE_WORDS or raw.endswith(_API_MEASURE_SUFFIXES)


def _measure_label(name: str, alone: bool) -> str:
    """A readable name for a measure column.

    The statistics API names its measures like OBSVALUE_OBSV; when a file
    has one measure, "القيمة" says as much and reads better.
    """
    raw = name.strip().upper()
    if not raw.endswith(_API_MEASURE_SUFFIXES):
        return name
    if alone:
        return "القيمة"
    for suffix in _API_MEASURE_SUFFIXES:
        raw = raw.removesuffix(suffix)
    return raw.replace("_", " ").strip().capitalize() or name


def _measures(columns: list[str], rows: list[tuple], skip: set[int]) -> list[int]:
    """Numeric columns, measure-named ones first."""
    found = [
        i for i in range(len(columns))
        if i not in skip and not _is_code(columns[i])
        and any(_is_number(r[i]) for r in rows)
        and all(r[i] is None or _is_number(r[i]) for r in rows)
    ]
    found.sort(key=lambda i: not _is_measure_name(columns[i]))
    return found[:MAX_CHARTS]


def _distinct(rows: list[tuple], columns: list[int]) -> set[tuple]:
    return {tuple(r[i] for i in columns) for r in rows}


def _over_years(columns, rows, year_at) -> list[dict]:
    measures = _measures(columns, rows, {year_at})
    others = [i for i in range(len(columns)) if i != year_at and i not in measures]
    if not measures or len(_distinct(rows, others)) > 1:
        return []
    years = [_year(r[year_at]) for r in rows]
    if None in years or len(set(years)) != len(years) or not 2 <= len(years) <= MAX_POINTS:
        return []
    charts = []
    for m in measures:
        points = sorted(
            ({"label": str(y), "value": r[m]} for y, r in zip(years, rows, strict=True)
             if _is_number(r[m])),
            key=lambda p: p["label"],
        )
        if len(points) >= 2:
            charts.append({"kind": "years", "note": "", "points": points,
                           "measure": _measure_label(columns[m], len(measures) == 1)})
    return charts


def _is_total(value: object) -> bool:
    return normalize(str(value or "")) in _TOTAL_WORDS


def _across(columns, rows, year_at) -> list[dict]:
    note = ""
    if year_at is not None:
        years = {_year(r[year_at]) for r in rows} - {None}
        if not years:
            return []
        latest = max(years)
        rows = [r for r in rows if _year(r[year_at]) == latest]
        note = f"سنة {latest}"
    skip = {year_at} if year_at is not None else set()
    measures = _measures(columns, rows, skip)
    varying = [
        i for i in range(len(columns))
        if i not in skip and i not in measures and len(_distinct(rows, [i])) > 1
    ]
    if not measures or not varying:
        return []
    # One breakdown: its label, English name and code change together.
    keys = _distinct(rows, varying)
    labels = [i for i in varying if len(_distinct(rows, [i])) == len(keys)]
    if not labels or len(keys) != len(rows):
        return []  # a second breakdown, or a category repeated
    label_at = next((i for i in labels if _ARABIC.search(str(rows[0][i] or ""))), labels[0])
    kept = [r for r in rows if not _is_total(r[label_at])]
    if len(kept) < len(rows):
        note = "، ".join(filter(None, (note, "دون صف الإجمالي")))
    if not 2 <= len(kept) <= MAX_CATEGORIES:
        return []
    charts = []
    for m in measures:
        points = sorted(
            ({"label": str(r[label_at]), "value": r[m]} for r in kept if _is_number(r[m])),
            key=lambda p: -p["value"],
        )
        if len(points) >= 2:
            charts.append({"kind": "categories", "note": note, "points": points,
                           "measure": _measure_label(columns[m], len(measures) == 1)})
    return charts


def _charts(columns: list[str], rows: list[tuple]) -> list[dict]:
    if len(rows) < 2:
        return []
    year_at = _year_column(columns)
    if year_at is not None:
        charts = _over_years(columns, rows, year_at)
        if charts:
            return charts
    return _across(columns, rows, year_at)


def read_preview(path: str | Path) -> dict | None:
    """The first rows of the workbook's data sheet, and a chart when one fits."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, BadZipFile):
        return None
    try:
        if DATA_SHEET not in workbook.sheetnames:
            return None
        rows = workbook[DATA_SHEET].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return None
        columns = [str(c) if c is not None else "" for c in header]
        while columns and not columns[-1]:
            columns.pop()
        width = len(columns)
        scanned: list[tuple] = []
        total = 0
        for row in rows:
            if row is None or all(v is None for v in row):
                continue
            total += 1
            if total <= MAX_SCAN_ROWS:
                scanned.append(tuple(row[:width]))
    finally:
        workbook.close()
    return {
        "columns": columns,
        "rows": [[_cell(v) for v in row] for row in scanned[:PREVIEW_ROWS]],
        "total_rows": total,
        "charts": _charts(columns, scanned) if total <= MAX_SCAN_ROWS else [],
    }
