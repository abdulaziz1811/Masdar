"""Reading published tables into a normalised form.

Official spreadsheets are messy in predictable ways, and three of those ways
decide whether the year a user asked for can be found at all:

* the year may be a column ("long" layout) or a set of column headers
  ("wide" layout, very common in statistical yearbooks);
* the real header row is often buried under title and logo rows;
* CSVs arrive with BOMs and with comma, semicolon or tab separators.

Getting the year axis right is what lets the verifier make a claim about a
year from the data itself instead of from a filename.
"""

from __future__ import annotations

import csv
import enum
import io
import json
import re
from dataclasses import dataclass, field

from masdar.nlu.normalize import contains_phrase, fold_digits, normalize, word_fit
from masdar.nlu.parser import extract_years

MAX_HEADER_SCAN = 15
MAX_ROWS = 200_000

_YEAR_COLUMN_NAMES = ("السنة", "سنة", "العام", "عام", "year", "السنه", "الفترة", "period")


# A contents sheet lists the workbook's tables -- every row naming the
# bulletin's year -- and is never the data. GASTAT's Hajj 2026 workbook opens
# with one, and it was once delivered as the answer's spreadsheet.
_INDEX_SHEETS = frozenset({"الفهرس", "فهرس", "المحتويات", "محتويات", "index", "contents", "toc"})
_INDEX_COLUMNS = ("رقم الجدول", "table no", "table number")

_INAPPLICABLE = frozenset({
    "لاينطبق", "لا ينطبق", "not-applicable", "not applicable", "not_applicable", "n/a",
})


class YearAxis(enum.Enum):
    ROWS = "rows"        # one column holds the year
    COLUMNS = "columns"  # each year is its own column
    NONE = "none"


@dataclass
class Table:
    columns: list[str]
    rows: list[list[object]]
    sheet_name: str = ""
    # Rows skipped above the header, kept so the export can show provenance
    # of the layout decision.
    preamble: list[str] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.rows)

    # -- layout --------------------------------------------------------
    def year_column_index(self) -> int | None:
        """Index of a column holding years, by name then by content."""
        for index, name in enumerate(self.columns):
            named_year = any(
                contains_phrase(str(name), candidate) for candidate in _YEAR_COLUMN_NAMES
            )
            if named_year and self._column_is_yearlike(index):
                return index
        for index in range(len(self.columns)):
            if self._column_is_yearlike(index, threshold=0.9):
                return index
        return None

    def _column_is_yearlike(self, index: int, threshold: float = 0.6) -> bool:
        values = [row[index] for row in self.rows[:500] if index < len(row)]
        values = [v for v in values if v not in (None, "")]
        if not values:
            return False
        hits = sum(1 for v in values if extract_years(str(v)))
        return hits / len(values) >= threshold

    def year_columns(self) -> dict[int, int]:
        """Map column index -> year, for headers that are themselves years."""
        found: dict[int, int] = {}
        for index, name in enumerate(self.columns):
            years = extract_years(str(name))
            # A header naming exactly one year is a year column; one naming a
            # range ("2015-2024") is a label, not a data column.
            if len(years) == 1:
                found[index] = years[0]
        return found

    def year_axis(self) -> YearAxis:
        year_cols = self.year_columns()
        # Years as headers only counts when most columns are years; a single
        # "2024" column next to labels is a long table with an odd name.
        if len(year_cols) >= 2 and len(year_cols) >= (len(self.columns) - len(year_cols)):
            return YearAxis.COLUMNS
        if self.year_column_index() is not None:
            return YearAxis.ROWS
        if year_cols:
            return YearAxis.COLUMNS
        return YearAxis.NONE

    # -- observation ---------------------------------------------------
    def observed_years(self) -> frozenset[int]:
        """Years actually present in the data. Authoritative evidence.

        A table with no year axis is often titled for its year inside the
        file («إجمالي أعداد الحجاج لعام 2026م» above a two-column table):
        that title is read from the file itself, and stands for the table.
        """
        axis = self.year_axis()
        if axis is YearAxis.COLUMNS:
            return frozenset(self.year_columns().values())
        if axis is YearAxis.ROWS:
            index = self.year_column_index()
            if index is None:
                return frozenset()
            years: set[int] = set()
            for row in self.rows:
                if index < len(row):
                    years.update(extract_years(str(row[index])))
            return frozenset(years)
        return self.title_years()

    def title_years(self) -> frozenset[int]:
        """Years named in the title rows above the header."""
        return frozenset(y for line in self.preamble for y in extract_years(line))

    @property
    def years_from_title(self) -> bool:
        return self.year_axis() is YearAxis.NONE and bool(self.title_years())

    @property
    def title(self) -> str:
        """What the table calls itself: its title rows, else its sheet's name."""
        return " ".join(self.preamble) or self.sheet_name

    @property
    def is_contents(self) -> bool:
        if normalize(self.sheet_name) in {normalize(n) for n in _INDEX_SHEETS}:
            return True
        header = " ".join(str(c) for c in self.columns)
        return any(contains_phrase(header, name) for name in _INDEX_COLUMNS)

    def observed_dimensions(self, dimension_words: dict) -> set:
        """Which breakdown axes the columns provide."""
        present = set()
        header_text = " ".join(str(c) for c in self.columns)
        for dimension, words in dimension_words.items():
            if any(contains_phrase(header_text, word) for word in words):
                present.add(dimension)
        return present

    # -- cleaning ------------------------------------------------------
    def without_inapplicable_columns(self) -> tuple[Table, list[str]]:
        """Drop columns whose every value says "not applicable".

        GASTAT's SDG tables return some forty columns per row, nearly all
        "لاينطبق" / "Not-Applicable". Removing a column that is constant and
        empty of meaning changes no figure, and the dropped names are
        returned so the workbook can list them.
        """
        if not self.rows:
            return self, []
        keep: list[int] = []
        dropped: list[str] = []
        for index, name in enumerate(self.columns):
            values = {
                str(row[index]).strip().lower() if index < len(row) and row[index] is not None
                else ""
                for row in self.rows
            }
            if values and values <= _INAPPLICABLE:
                dropped.append(str(name))
            else:
                keep.append(index)
        if not dropped:
            return self, []
        columns = [self.columns[i] for i in keep]
        rows = [[row[i] if i < len(row) else None for i in keep] for row in self.rows]
        return Table(columns, rows, self.sheet_name, list(self.preamble)), dropped

    # -- slicing -------------------------------------------------------
    def filter_places(
        self, wanted: dict[str, tuple[str, ...]], known: dict[str, tuple[str, ...]]
    ) -> tuple[Table, bool]:
        """Rows for the wanted regions, and whether the table could be narrowed.

        Only a column that actually lists regions is used -- at least three
        rows naming a known region -- so «الرياض» cannot match «رياض الأطفال»
        in a column about school stages. A table without such a column, or
        without a row for the region, is returned whole.
        """
        from masdar.nlu.normalize import contains_phrase

        def names(cell: object, places: dict[str, tuple[str, ...]]) -> bool:
            text = str(cell or "")
            return bool(text) and any(
                contains_phrase(text, v) for variants in places.values() for v in variants
            )

        width = max((len(r) for r in self.rows), default=0)
        for index in range(width):
            cells = [row[index] if index < len(row) else None for row in self.rows]
            if sum(1 for c in cells if names(c, known)) < 3:
                continue
            kept = [row for row, cell in zip(self.rows, cells, strict=True) if names(cell, wanted)]
            if kept:
                return Table(list(self.columns), kept, self.sheet_name, list(self.preamble)), True
            return self, False
        return self, False

    def filter_years(self, years: frozenset[int]) -> Table:
        """Narrow the table to the requested years, whatever the layout."""
        if not years:
            return self
        axis = self.year_axis()

        if axis is YearAxis.ROWS:
            index = self.year_column_index()
            if index is None:
                return self
            kept = [
                row for row in self.rows
                if index < len(row) and years & set(extract_years(str(row[index])))
            ]
            return Table(list(self.columns), kept, self.sheet_name, list(self.preamble))

        if axis is YearAxis.COLUMNS:
            year_cols = self.year_columns()
            keep_indexes = [i for i in range(len(self.columns)) if i not in year_cols]
            keep_indexes += [i for i, y in sorted(year_cols.items()) if y in years]
            keep_indexes.sort()
            columns = [self.columns[i] for i in keep_indexes]
            rows = [[row[i] if i < len(row) else None for i in keep_indexes] for row in self.rows]
            return Table(columns, rows, self.sheet_name, list(self.preamble))

        return self


_NUMERIC = re.compile(r"^-?\d{1,3}(?:,\d{3})*(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")


def coerce(value: object) -> object:
    """Turn numeric-looking text into a number.

    CSV gives everything back as strings; leaving them that way produces a
    spreadsheet whose columns cannot be summed or charted, which defeats the
    point of exporting one. Percentages, codes and anything ambiguous are
    left as text.
    """
    if not isinstance(value, str):
        return value
    text = fold_digits(value.strip())
    if not text or not _NUMERIC.match(text):
        return value.strip()
    bare = text.replace(",", "")
    # A leading zero is usually an identifier, not a quantity.
    if len(bare.lstrip("-")) > 1 and bare.lstrip("-").startswith("0") and "." not in bare:
        return value.strip()
    try:
        number = float(bare)
    except ValueError:
        return value.strip()
    return int(number) if number.is_integer() and "." not in bare else number


def _is_number(cell: object) -> bool:
    if isinstance(cell, bool):
        return False
    if isinstance(cell, int | float):
        return True
    return bool(_NUMERIC.match(fold_digits(str(cell)).strip()))


def _find_header_row(grid: list[list[object]]) -> int:
    """Pick the header row, skipping title and logo rows above it.

    The header is the row with the most filled cells among the first few,
    provided real rows follow it; ties go to the earliest row. A row holding
    figures is data, and is a header only when no row without figures will
    do: GASTAT's gender table (ذكور | إناث above 893396 | 813905 | 1707301)
    once had its numbers read as column names, and a yearbook's «— | 2019 |
    2020» lost to the fuller «الرياض | 10 | 11» below it. Years are words here.
    """
    def figures(row: list[object]) -> bool:
        return any(
            _is_number(c) and not extract_years(str(c)) for c in row if c not in (None, "")
        )

    def pick(allow_figures: bool) -> int | None:
        best_index, best_filled = None, 1
        for index in range(min(MAX_HEADER_SCAN, len(grid))):
            filled = sum(1 for cell in grid[index] if cell not in (None, ""))
            if filled <= best_filled or index + 1 >= len(grid):
                continue
            if not allow_figures and figures(grid[index]):
                continue
            best_index, best_filled = index, filled
        return best_index

    found = pick(allow_figures=False)
    if found is None:
        found = pick(allow_figures=True)
    return found or 0


def _grid_to_table(grid: list[list[object]], sheet_name: str = "") -> Table:
    grid = [row for row in grid if any(cell not in (None, "") for cell in row)]
    if not grid:
        return Table([], [], sheet_name)
    header_index = _find_header_row(grid)
    header = grid[header_index]
    columns = [
        (str(cell).strip() if cell not in (None, "") else f"عمود {i + 1}")
        for i, cell in enumerate(header)
    ]
    rows = [row for row in grid[header_index + 1:]]
    preamble = [
        " | ".join(str(c) for c in row if c not in (None, ""))
        for row in grid[:header_index]
    ]
    return Table(columns, rows, sheet_name, preamble)


def read_csv(content: bytes) -> Table:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid = [
        [coerce(cell) for cell in row]
        for _, row in zip(range(MAX_ROWS), reader, strict=False)
    ]
    return _grid_to_table(grid)


def read_xlsx(content: bytes, question: str = "") -> Table:
    """The workbook's sheet that best answers `question`.

    A bulletin's workbook holds one table per sheet («إجمالي أعداد الحجاج»,
    «... حسب الجنس», «... حسب طريقة القدوم»): the one worded most like the
    question is the answer, not merely the longest. A contents sheet never is.
    """
    from openpyxl import load_workbook

    from masdar.nlu.lexicon import load_lexicon

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    stopwords = load_lexicon().stopwords if question else frozenset()

    def rank(table: Table) -> tuple[bool, float, bool, int]:
        # Then a sheet carrying years beats one without, and among equals the
        # longer one wins. Yearbooks often open with a contents or notes sheet.
        fit = word_fit(question, table.title, stopwords) if question else 0.0
        return not table.is_contents, round(fit, 3), bool(table.observed_years()), len(table)

    best: Table | None = None
    for sheet in workbook.worksheets:
        grid = [
            list(row)
            for _, row in zip(range(MAX_ROWS), sheet.iter_rows(values_only=True), strict=False)
        ]
        table = _grid_to_table(grid, sheet.title)
        if best is None or rank(table) > rank(best):
            best = table
    workbook.close()
    return best or Table([], [])


# OData bookkeeping keys carry no data.
_ODATA = ("@odata", "odata.")

# GASTAT's chart endpoint names its columns per indicator, e.g. POP_TIME and
# POP_OBSV, so the period and value columns can only be found by suffix.
_TIME_SUFFIX = re.compile(r"_TIME$", re.IGNORECASE)
_VALUE_SUFFIX = re.compile(r"_OBSV$", re.IGNORECASE)


def _json_rows(payload: object) -> list[dict]:
    """Find the list of records in a JSON body.

    Bodies seen in the wild: a bare array, an OData `{"value": [...]}`, and
    assorted `{"data": [...]}` wrappers. All three are accepted rather than
    assuming one.
    """
    if isinstance(payload, list):
        return [r for r in payload if isinstance(r, dict)]
    if isinstance(payload, dict):
        for key in ("value", "data", "info", "results", "items", "records"):
            candidate = payload.get(key)
            if isinstance(candidate, list):
                return [r for r in candidate if isinstance(r, dict)]
    return []


def read_json(content: bytes) -> Table:
    """Read a JSON array of flat records into a table.

    Column order follows first appearance so a time column stays leftmost,
    which is also the order the year-axis detection expects.
    """
    try:
        payload = json.loads(content.decode("utf-8", errors="replace"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON: {exc}") from exc

    records = _json_rows(payload)
    if not records:
        return Table([], [])

    columns: list[str] = []
    for record in records:
        for key in record:
            if any(key.lower().startswith(prefix) for prefix in _ODATA):
                continue
            if key not in columns:
                columns.append(key)

    # Put the period column first; everything downstream reads more naturally.
    timed = [c for c in columns if _TIME_SUFFIX.search(c)]
    if timed:
        columns = timed + [c for c in columns if c not in timed]

    rows = [[coerce_value(record.get(column)) for column in columns] for record in records]
    return Table(columns, rows)


def coerce_value(value: object) -> object:
    return coerce(value) if isinstance(value, str) else value


def read_table(content: bytes, fmt: str, question: str = "") -> Table:
    fmt = (fmt or "").upper()
    if fmt in ("CSV", "TSV", "TXT"):
        return read_csv(content)
    if fmt in ("XLSX", "XLSM", "XLS"):
        return read_xlsx(content, question)
    if fmt == "JSON":
        return read_json(content)
    raise ValueError(f"unsupported tabular format: {fmt or 'unknown'}")
