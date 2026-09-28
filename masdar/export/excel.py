"""Workbook export.

A spreadsheet outlives the conversation that produced it, so every workbook
carries a "المصدر" sheet recording where the numbers came from, when they
were retrieved, the publisher's own last-update date and the file's sha256.
A file that leaves here without that sheet is an unattributed number, which
is exactly what this project exists to avoid.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from masdar.domain.models import Coverage, DataRequest, DatasetCandidate, Provenance, Verdict
from masdar.export.tabular import Table

DATA_SHEET = "البيانات"
SOURCE_SHEET = "المصدر"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LABEL_FONT = Font(bold=True)
_WARNING_FILL = PatternFill("solid", fgColor="C00000")

MAX_COLUMN_WIDTH = 55
_UNSAFE = re.compile(r"[^\w؀-ۿ-]+")

VERDICT_LABELS = {
    Verdict.AVAILABLE: "متوفرة ومؤكَّدة",
    Verdict.PARTIAL: "متوفرة جزئياً",
    Verdict.NOT_AVAILABLE: "غير متوفرة",
    Verdict.UNVERIFIED: "غير مؤكَّدة",
    Verdict.NO_SOURCE: "لا يوجد مصدر",
    Verdict.SOURCE_UNREACHABLE: "تعذّر الوصول إلى المصدر",
}

SYNTHETIC_WARNING = (
    "تحذير: هذه بيانات تجريبية مُولّدة لاختبار النظام وليست بيانات رسمية. "
    "لا تُستخدم كمرجع."
)


def safe_filename(*parts: str) -> str:
    cleaned = [_UNSAFE.sub("_", p.strip()).strip("_") for p in parts if p and p.strip()]
    return "_".join(filter(None, cleaned)) or "masdar"


def _autosize(sheet, widths: dict[int, int]) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)


def _write_data_sheet(sheet, table: Table) -> None:
    sheet.sheet_view.rightToLeft = True
    widths: dict[int, int] = {}

    for column, name in enumerate(table.columns, start=1):
        cell = sheet.cell(row=1, column=column, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths[column] = len(str(name))

    for row_index, row in enumerate(table.rows, start=2):
        for column, value in enumerate(row[: len(table.columns)], start=1):
            sheet.cell(row=row_index, column=column, value=value)
            widths[column] = max(widths.get(column, 0), len(str(value)) if value is not None else 0)

    _autosize(sheet, widths)
    sheet.freeze_panes = "A2"


def _write_source_sheet(
    sheet,
    request: DataRequest,
    candidate: DatasetCandidate,
    provenance: Provenance,
    verdict: Verdict,
    coverage: Coverage,
    notes: tuple[str, ...],
    synthetic: bool,
    row_count: int,
) -> None:
    sheet.sheet_view.rightToLeft = True
    row = 1

    if synthetic:
        cell = sheet.cell(row=row, column=1, value=SYNTHETIC_WARNING)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _WARNING_FILL
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 2

    def field(label: str, value: object) -> None:
        nonlocal row
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = _LABEL_FONT
        label_cell.alignment = Alignment(vertical="top")
        text = "" if value in (None, "") else str(value)
        value_cell = sheet.cell(row=row, column=2, value=text)
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    def heading(text: str) -> None:
        nonlocal row
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1

    heading("الطلب")
    field("السؤال كما ورد", request.raw_query)
    field("الموضوع", request.topic.label_ar if request.topic else "غير محدد")
    field("الفترة المطلوبة", request.period.label())
    if request.dimensions:
        field("التفصيل المطلوب", "، ".join(d.label_ar for d in request.dimensions))
    row += 1

    heading("النتيجة")
    field("حالة التوفر", VERDICT_LABELS.get(verdict, verdict.value))
    field("عدد الصفوف في هذا الملف", row_count)
    field(
        "السنوات المتاحة في المصدر",
        "، ".join(str(y) for y in sorted(coverage.years)) or "غير معروفة",
    )
    if coverage.years:
        field("نطاق التغطية", coverage.describe())
    field("أساس التحقق", COVERAGE_ORIGIN_LABELS.get(coverage.origin.value, coverage.origin.value))
    row += 1

    heading("المصدر")
    field("الجهة الناشرة", provenance.publisher_ar or candidate.publisher_ar)
    field("اسم مجموعة البيانات", candidate.title_ar or candidate.title_en)
    field("صفحة المصدر", provenance.landing_url)
    field("رابط الملف", provenance.resource_url or "—")
    field(
        "آخر تحديث معلن من المصدر",
        provenance.last_updated.isoformat() if provenance.last_updated else "غير معلن",
    )
    if provenance.last_updated_note:
        field("ملاحظة على تاريخ التحديث", provenance.last_updated_note)
    field("تاريخ الاستخراج", provenance.retrieved_at.isoformat(timespec="seconds"))
    if provenance.stale:
        field("تنبيه", "نسخة محفوظة: تعذّر الوصول إلى المصدر وقت الطلب، وقد توجد بيانات أحدث")
    field("الترخيص", provenance.license_name or "غير محدد")
    field("بصمة الملف (SHA-256)", provenance.sha256 or "—")
    field("حجم الملف (بايت)", provenance.byte_size if provenance.byte_size is not None else "—")
    row += 1

    if notes:
        heading("ملاحظات")
        for note in notes:
            field("•", note)

    _autosize(sheet, {1: 28, 2: 70})


COVERAGE_ORIGIN_LABELS = {
    "observed_data": "تم فتح الملف والتحقق من السنوات داخله",
    "metadata_claim": "بيانات وصفية معلنة من المصدر",
    "inferred_title": "مستنتج من العنوان (غير مؤكد)",
}


def write_workbook(
    path: Path,
    table: Table,
    request: DataRequest,
    candidate: DatasetCandidate,
    provenance: Provenance,
    verdict: Verdict,
    coverage: Coverage,
    notes: tuple[str, ...] = (),
    synthetic: bool = False,
) -> Path:
    """Write the data and its provenance to `path`, returning the path."""
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = DATA_SHEET
    _write_data_sheet(data_sheet, table)

    source_sheet = workbook.create_sheet(SOURCE_SHEET)
    _write_source_sheet(
        source_sheet,
        request,
        candidate,
        provenance,
        verdict,
        coverage,
        notes,
        synthetic,
        len(table.rows),
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def default_path(
    out_dir: Path,
    request: DataRequest,
    candidate: DatasetCandidate,
    years: tuple[int, ...] = (),
    now: datetime | None = None,
) -> Path:
    """Name the file after the years it actually contains.

    Naming it after the years that were *requested* would label a partial
    result as though it were complete.
    """
    now = now or datetime.now()
    topic = request.topic.id if request.topic else "data"
    if years:
        period = str(years[0]) if len(years) == 1 else f"{min(years)}-{max(years)}"
    else:
        period = request.period.label().replace("–", "-").replace(" ", "")
    # The dataset id keeps two datasets of one source, exported in the same
    # second for the same question, from overwriting each other.
    name = safe_filename(
        "masdar", topic, period, candidate.source_id, candidate.dataset_id[-24:],
        now.strftime("%Y%m%d-%H%M%S"),
    )
    return out_dir / f"{name}.xlsx"
